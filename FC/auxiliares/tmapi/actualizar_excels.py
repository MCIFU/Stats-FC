"""Recalcula con Transfermarkt (partido a partido) las columnas LIGA/COPA/CONT/FIFA/SEL de TODAS las filas de
Temporada 2025-26.xlsx y Temporada 2026-27.xlsx (método de LEEME_METODO_TM_API.md) y compara con lo que hay.

Uso (desde la carpeta FC):
    python auxiliares/tmapi/actualizar_excels.py                 # descarga + informe de diferencias (no escribe)
    python auxiliares/tmapi/actualizar_excels.py --escribir      # además escribe las celdas que cambian
    python auxiliares/tmapi/actualizar_excels.py --corte 2026-09-24

No toca filas, orden, Club, Liga, NAC, POS, Edad ni Valor: solo PJ/Min/G/A (GC/CS en porteros) de cada bloque y el
código CONT. Filas sin ID TM se dejan como están. Informe: auxiliares/tmapi/diferencias_TM.csv
"""
import argparse, collections, concurrent.futures as cf, copy, datetime as dt, json, sys, time
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import descargar_partidos as D  # noqa: E402  (descarga, caché, cruce de IDs y lectura de performance-game)

FC = D.FC
SHEETS = D.SHEETS
FILES = {"2025-26": "Temporada 2025-26.xlsx", "2026-27": "Temporada 2026-27.xlsx"}
CAL_LIGAS = {"MLS", "Brasileirão", "Liga Argentina"}  # ligas de año natural: temporada TM = año - 1
CAL_COMPS = {"MLS1", "USL", "BRA1", "BRA2", "ARG1", "AR1N", "ARGC"}  # ligas cuyo club juega por año natural
CAL_CUPS = {"USMX", "CAMC", "CCL", "CCAC", "CCLS"}  # Leagues Cup, Campeones Cup, Concachampions: temporada TM = año
BLOCKS = ["LIGA", "COPA", "CONT", "FIFA", "SEL"]
AMBIGUOS = {("Athletic Club", "Campeonato Brasileiro Série B")}  # clubmap solo tiene el Athletic de Bilbao
COLS = {"LIGA": 7, "COPA": 11, "CONT": 16, "FIFA": 20, "SEL": 24}  # 1ª columna (PJ) de cada bloque; CONT. = col 15
CONT_CODE = {'CL': 'UCL', 'CLQ': 'UCL', 'USC': 'UCL', 'EL': 'UEL', 'ELQ': 'UEL', 'UCOL': 'UECL', 'ECLQ': 'UECL',
             'CLI': 'LIB', 'CLIQ': 'LIB', 'CS': 'SUD', 'RECO': 'LIB', 'CCL': 'CCC', 'CCAC': 'CCC', 'CCLS': 'CCC',
             'USMX': 'LC', 'CAMC': 'CC', 'ACLE': 'ACL', 'ACEQ': 'ACL', 'ACL2': 'ACL', 'AC2Q': 'ACL', 'AFCP': 'ACL',
             'CAFC': 'CAF', 'CAFCC': 'CAF'}
PRIO = ['UCL', 'UEL', 'UECL', 'LIB', 'SUD', 'CCC', 'LC', 'CC', 'ACL', 'CAF']


def excel_rows():
    cm = json.loads((HERE / "clubmap.json").read_text(encoding="utf-8"))
    # clubes que no están en clubmap.json (filas creadas con el nombre TM): nombre o nombre corto TM exacto y único
    mc = D.CACHE / "meta_clubs.json.gz"
    if mc.exists():
        import gzip
        by_name = collections.defaultdict(set)
        for k, v in json.loads(gzip.decompress(mc.read_bytes())).items():
            b = v.get("baseDetails") or {}
            if b.get("isNationalTeam") or str(b.get("mainClubId") or k) != k:
                continue  # ni selecciones ni filiales
            for n in {v.get("name"), b.get("shortName")} - {None}:
                by_name[n].add(k)
        extra = {n: next(iter(ids)) for n, ids in by_name.items() if len(ids) == 1 and n not in cm}
        cm = {**extra, **cm}
    ce = HERE / "clubmap_extra.json"  # clubes añadidos por ampliar_ligas.py (nombre corto TM resuelto dentro de su liga)
    if ce.exists():
        cm = {**cm, **{k: v for k, v in json.loads(ce.read_text(encoding="utf-8")).items() if k not in cm}}
    out = []
    for season, fname in FILES.items():
        wb = openpyxl.load_workbook(FC / fname, read_only=True)
        for sh in SHEETS:
            rows = list(wb[sh].iter_rows(values_only=True))
            hi = next(i for i, r in enumerate(rows[:10]) if r and r[0] == "Jugador")
            for k, r in enumerate(rows[hi + 1:], hi + 2):
                if not r or not r[0] or not r[4]:
                    continue
                tms = (2024 if r[5] in CAL_LIGAS else 2025) if season == "2025-26" else (2025 if r[5] in CAL_LIGAS else 2026)
                cid = str(cm[r[4]]) if r[4] in cm else None
                if (r[4], r[5]) in AMBIGUOS:
                    cid = None  # mismo nombre que otro club del mapa: no se puede enlazar con seguridad
                out.append(dict(season=season, tmseason=tms, sheet=sh, row=k, name=r[0], nat=r[3], club=r[4],
                                club_id=cid, liga=r[5],
                                old=list(r[6:27]), gk=sh == "PORTEROS"))
    return out


def construir(pids, offline=False, workers=6, corte=None, extra_clubs=()):
    """Descarga (o lee de caché) performance-game de `pids` y devuelve el contexto: partidos válidos por jugador con su
    temporada/bloque/confederación, metadatos de clubes y competiciones."""
    a = argparse.Namespace(offline=offline, workers=workers, corte=corte or (dt.date.today() - dt.timedelta(days=1)).isoformat())
    rows = [dict(club_id=c) for c in extra_clubs]
    perf, t0 = {}, time.time()

    def compact(p):
        """Solo partidos desde 2025 y los campos que se usan (la carrera entera no cabe en memoria para 7.000 jugadores)."""
        out = []
        for e in D.performance(p, a.offline):
            if str(D.dig(e, "gameInformation", "date", "dateTimeUTC") or "") < "2025-01-01":
                continue
            out.append((D.parse_game(e), D.dig(e, "clubsInformation", "club", "opponentGoalsTotal")))
        return out
    with cf.ThreadPoolExecutor(a.workers) as ex:
        futs = {ex.submit(compact, p): p for p in pids}
        for i, f in enumerate(cf.as_completed(futs), 1):
            try:
                perf[futs[f]] = f.result()
            except Exception as e:  # noqa: BLE001
                print("  !", futs[f], e)
            if i % 1000 == 0:
                print(f"  {i}/{len(pids)} ({time.time() - t0:.0f}s)")
    games = perf
    comps = {g["comp"] for L in games.values() for g, _ in L}
    clubs = {g["club_id"] for L in games.values() for g, _ in L} | {r["club_id"] for r in rows}
    CM = D.meta("competitions", comps, a.offline)
    KM = D.meta("clubs", clubs, a.offline)
    COMP = {k: {"t": v.get("typeId"), "conf": D.dig(v, "originDetails", "confederationId"),
                "ctry": D.dig(v, "originDetails", "countryId")} for k, v in CM.items()}
    ctry_conf = collections.Counter()
    for c in COMP.values():
        if c["ctry"] and c["conf"]:
            ctry_conf[(c["ctry"], c["conf"])] += 1
    CTRY_CONF = {}
    for (ct, co), n in sorted(ctry_conf.items(), key=lambda x: -x[1]):
        CTRY_CONF.setdefault(ct, co)
    CLUB = {}
    for k, v in KM.items():
        b = v.get("baseDetails") or {}
        pc = b.get("primaryCompetitionId")
        conf = COMP.get(pc, {}).get("conf") or CTRY_CONF.get(b.get("countryId"))
        CLUB[k] = {"nt": bool(b.get("isNationalTeam")), "main": str(b.get("mainClubId") or k), "conf": conf,
                   "ctry": b.get("countryId"), "name": b.get("shortName") or v.get("name")}
    cal_ctry = {COMP.get(c, {}).get("ctry") for c in ("MLS1", "BRA1", "ARG1", "AR1N", "JAP1")} - {None}  # JAP1: año natural hasta 2025
    # liga de año natural = su competición principal va una temporada por detrás de LaLiga en TM (sept-2026: 2025 vs 2026)
    prim = {(v.get("baseDetails") or {}).get("primaryCompetitionId") for v in KM.values()} - {None}
    CM.update(D.meta("competitions", prim, a.offline))
    ref = (CM.get("ES1") or {}).get("currentSeasonId")
    cal_comp = {c for c in prim if ref and (CM.get(c) or {}).get("currentSeasonId") == ref - 1} | CAL_COMPS
    cal_clubs = {k for k, v in KM.items() if (v.get("baseDetails") or {}).get("primaryCompetitionId") in cal_comp
                 or CLUB[k]["ctry"] in cal_ctry}  # MLS incluye clubes canadienses
    print("clubes de año natural:", len(cal_clubs), "| ligas:", sorted(cal_comp)[:40])
    D_COMP = {k: {"t": v["t"]} for k, v in COMP.items()}
    cut = pd.Timestamp(a.corte + " 20:55", tz="UTC")

    def tag_of(g, blk):
        d = pd.Timestamp(g["date"])
        d = d.tz_localize("UTC") if d.tzinfo is None else d.tz_convert("UTC")
        if blk == "SEL":  # selección: agosto-julio por fecha
            return "2025-26" if pd.Timestamp("2025-08-01", tz="UTC") <= d < pd.Timestamp("2026-08-01", tz="UTC") else \
                "2026-27" if d >= pd.Timestamp("2026-08-01", tz="UTC") else None
        if str(g["comp"]) in CAL_CUPS and g["club_id"] not in cal_clubs:  # copas CONCACAF de año natural en clubes MX: temporada julio-junio por fecha
            return "2025-26" if pd.Timestamp("2025-07-01", tz="UTC") <= d < pd.Timestamp("2026-07-01", tz="UTC") else \
                "2026-27" if d >= pd.Timestamp("2026-07-01", tz="UTC") else None
        if g["club_id"] in cal_clubs:  # año natural: por año de la fecha
            return {2025: "2025-26", 2026: "2026-27"}.get(d.year)
        return {2025: "2025-26", 2026: "2026-27"}.get(int(g["season_id"] or 0))

    # partidos válidos por jugador: (temporada, bloque, club, conf, fecha, valores)
    valid = collections.defaultdict(list)
    for pid, L in games.items():
        for g, e in L:
            played = (g["Min"] or 0) > 0 or g["state"] == "played"
            if not played or g["live"] or not g["date"]:
                continue
            d = pd.Timestamp(g["date"])
            d = d.tz_localize("UTC") if d.tzinfo is None else d.tz_convert("UTC")
            if d > cut:
                continue
            blk = D.block_of(g, D_COMP, CLUB)
            if blk is None:
                continue
            if g["comp"] == "KLUB" and str(g["season_id"]) == "2024":
                continue
            tag = tag_of(g, blk)
            if tag is None:
                continue
            opp_goals = e
            v = [1, g["Min"] or 0, g["G"] or 0, g["A"] or 0, g["GC"] or 0, 1 if opp_goals == 0 else 0]
            valid[pid].append(dict(tag=tag, blk=blk, club=g["club_id"], conf=CLUB.get(g["club_id"], {}).get("conf"),
                                   comp=str(g["comp"]), date=d, v=v, g=g))

    return dict(games=games, valid=valid, CLUB=CLUB, COMP=COMP, CM=CM, KM=KM, cal_clubs=cal_clubs)


def agregar(rows, ctx):
    """Bloques LIGA/COPA/CONT/FIFA/SEL de cada fila (reglas del método) + partidos usados. rows: dicts con tm_id."""
    valid, CLUB = ctx["valid"], ctx["CLUB"]
    by_pid = collections.defaultdict(list)
    for r in rows:
        if r.get("tm_id"):
            by_pid[(r["tm_id"], r["season"])].append(r)
    res, stats = [], collections.Counter()
    for (pid, season), L in by_pid.items():
        G = [x for x in valid.get(pid, []) if x["tag"] == season]
        clubg = [x for x in G if x["blk"] != "SEL"]
        fin = pd.Timestamp("2026-08-01", tz="UTC") if season == "2025-26" else pd.Timestamp("2100-01-01", tz="UTC")
        clubg_fin = [x for x in clubg if x["date"] < fin] or clubg  # club al final de la temporada (antes de agosto)
        last = max(clubg_fin, key=lambda x: x["date"])["conf"] if clubg else None
        confs = [CLUB.get(r["club_id"], {}).get("conf") for r in L]
        # SEL solo en la fila del continente donde acabó la temporada; si esa fila no está en el Excel, en ninguna
        sel_owner = next((i for i, c in enumerate(confs) if c == last), -1) if last else 0
        for i, r in enumerate(L):
            conf = confs[i]
            out = {b: [0] * 6 for b in BLOCKS}
            used = []
            cont = collections.Counter()
            for x in G:
                if x["blk"] == "SEL":
                    if i != sel_owner:
                        continue
                elif conf is not None and x["conf"] != conf:
                    continue  # regla de continentes: cada fila solo con clubes de su confederación
                for j in range(6):
                    out[x["blk"]][j] += x["v"][j]
                used.append(x)
                if x["blk"] == "CONT":
                    cont[CONT_CODE.get(x["comp"], x["comp"])] += 1
            idx = (0, 1, 4, 5) if r["gk"] else (0, 1, 2, 3)
            new = []
            for b in BLOCKS:
                new += [out[b][j] for j in idx]
                if b == "COPA":
                    code = None
                    if cont:
                        mx = max(cont.values())
                        code = sorted([c for c, n in cont.items() if n == mx], key=lambda c: PRIO.index(c) if c in PRIO else 99)[0]
                    new.append(code)
            old = r.get("old") or [0] * 8 + [None] + [0] * 12
            num_old = [x if isinstance(x, (int, float)) else 0 for i2, x in enumerate(old) if i2 != 8]
            num_new = [x for i2, x in enumerate(new) if i2 != 8]
            diff = [BLOCKS[k // 4] + " " + ["PJ", "Min", "G/GC", "A/CS"][k % 4] for k in range(20) if num_old[k] != num_new[k]]
            cont_old = old[8] if old[8] not in (None, "", "—", "-") else None
            cont_diff = (new[8] or None) != cont_old and num_new[8] > 0  # solo si ahora hay partidos continentales
            stats["cambian" if diff else "iguales"] += 1
            res.append(dict(season=season, sheet=r["sheet"], row=r["row"], name=r["name"], club=r["club"], liga=r["liga"],
                            tm_id=pid, cambios=", ".join(diff), cont_old=old[8], cont_new=new[8], cont_cambia=bool(cont_diff),
                            old=old, new=new, used=used, excel_club=r["club"], gk=r["gk"]))
    return res, stats


def exportar_partidos(res, ligas, out_path):
    """Una fila por jugador y partido (sin selección) para el motor: ELO/FORM/CONSISTENCY/OPPONENT_STRENGTH."""
    out = []
    for r in res:
        if ligas and r["liga"] not in ligas:
            continue
        for x in r["used"]:
            if x["blk"] == "SEL":
                continue
            g = x["g"]
            out.append(dict(season=r["season"], source_sheet=r["sheet"], name=r["name"], excel_club=r["excel_club"], liga=r["liga"],
                            tm_id=r["tm_id"], date=x["date"].strftime("%Y-%m-%d %H:%M"), block=x["blk"], comp=x["comp"],
                            game_id=g["game_id"], club_id=g["club_id"], club=CTX_CLUB.get(g["club_id"], {}).get("name"),
                            opponent_id=g["opp_id"], opponent=CTX_CLUB.get(g["opp_id"] or "", {}).get("name") or g["opp_id"],
                            home=g["home"], gf=g["gf"], ga=g["ga"], Min=g["Min"], G=g["G"], A=g["A"],
                            GC=g["GC"] if r["gk"] else None))
    df = pd.DataFrame(out).sort_values(["season", "date", "tm_id"])
    df.to_csv(out_path, index=False, encoding="utf-8-sig")
    print(f"{len(df)} filas jugador-partido ({df.name.nunique()} jugadores) -> {out_path}")
    return df


CTX_CLUB = {}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corte", default=(dt.date.today() - dt.timedelta(days=1)).isoformat())
    ap.add_argument("--escribir", action="store_true")
    ap.add_argument("--filas", help="solo estas filas 'temporada|HOJA|fila' separadas por comas (las revisadas a mano)")
    ap.add_argument("--nuevas", action="store_true", help="con --escribir: solo filas sin ningún dato todavía (añadidas por ampliar_ligas.py)")
    ap.add_argument("--partidos", help="ligas (separadas por comas, o 'todas') para exportar partidos_TM.csv")
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--crecimiento", action="store_true",
                    help="con --escribir: solo filas cuyos números solo suben (partidos nuevos); las que bajan quedan en el informe")
    a = ap.parse_args()
    rows = excel_rows()
    sin_club = [r for r in rows if r["club_id"] is None]
    print(f"{len(rows)} filas; {len(sin_club)} con club fuera de clubmap.json (se dejan igual)")
    rows = D.map_ids([r for r in rows if r["club_id"]])
    pids = sorted({r["tm_id"] for r in rows if r.get("tm_id")})
    print(f"{len(pids)} jugadores TM; performance-game…")
    ctx = construir(pids, a.offline, a.workers, a.corte, {r["club_id"] for r in rows})
    CTX_CLUB.update(ctx["CLUB"])
    res, stats = agregar(rows, ctx)
    if a.partidos:
        ligas = None if a.partidos == "todas" else set(a.partidos.split(","))
        exportar_partidos(res, ligas, FC / "Claude outputs" / "partidos_TM.csv")
    res = [{k: v for k, v in r.items() if k != "used"} for r in res]
    df = pd.DataFrame(res)
    rep = df[(df.cambios != "") | df.cont_cambia].copy()
    rep["antes"] = rep.old.map(lambda o: " | ".join("/".join(str(x) for x in o[i:i + 4]) for i in (0, 4, 9, 13, 17)))
    rep["ahora"] = rep.new.map(lambda o: " | ".join("/".join(str(x) for x in o[i:i + 4]) for i in (0, 4, 9, 13, 17)))
    rep.drop(columns=["old", "new"]).to_csv(HERE / "diferencias_TM.csv", index=False, encoding="utf-8-sig")
    print("filas:", dict(stats), "| cambia CONT:", int(df.cont_cambia.sum()))
    print("bloques que cambian:", collections.Counter(c.split()[0] for s in rep.cambios for c in s.split(", ") if c).most_common())
    print("por temporada:", rep.groupby("season").size().to_dict(), "-> auxiliares/tmapi/diferencias_TM.csv")
    df.to_pickle(HERE / "cache" / "recalculo.pkl")
    if a.escribir:
        if a.filas:
            ok = set(a.filas.split(","))
            df = df[(df.season + "|" + df.sheet + "|" + df.row.astype(str)).isin(ok)]
            print(f"escribiendo solo {len(df)} filas revisadas")
        if a.crecimiento:
            num = lambda x: x if isinstance(x, (int, float)) else 0
            sube = df.apply(lambda r: all(num(n) >= num(o) for n, o in zip(r.new, r.old)), axis=1)
            print(f"escribiendo {int(sube.sum())} filas que solo crecen; {int((~sube & (df.cambios != '')).sum())} con bajadas quedan en diferencias_TM.csv")
            df = df[sube]
        if a.nuevas:
            df = df[df.old.map(lambda o: all(not isinstance(x, (int, float)) or x == 0 for x in o))]
            print(f"escribiendo solo {len(df)} filas nuevas (sin datos previos)")
        escribir(df, a.corte)




def escribir(df, corte):
    for season, fname in FILES.items():
        wb = openpyxl.load_workbook(FC / fname)
        n = 0
        for r in df[df.season == season].itertuples():
            if not r.cambios and not r.cont_cambia:
                continue
            ws = wb[r.sheet]
            for k, v in enumerate(r.new):
                col = 7 + k
                if col == 15:
                    if r.cont_cambia:
                        ws.cell(r.row, col).value = v or "—"
                    continue
                ws.cell(r.row, col).value = v
            n += 1
        rebuild_rankings(wb)
        ws = [wb[s] for s in wb.sheetnames if s.startswith("LEEME")][0]
        c = ws.cell(ws.max_row + 1, 1)
        c.value = (f"ACT {dt.date.today():%d/%m/%Y} (Claude): recálculo Transfermarkt partido a partido de todas las filas "
                   f"(corte: saque inicial ≤ 20:55 UTC del {corte}). {n} filas corregidas; detalle en auxiliares/tmapi/diferencias_TM.csv. "
                   "Filas, orden, Club/Liga/NAC/POS/Edad/Valor sin tocar.")
        c._style = copy.copy(ws.cell(ws.max_row - 1, 1)._style)
        wb.save(FC / fname)
        print(f"{fname}: {n} filas escritas")


def rebuild_rankings(wb):
    """RANKING_TOTAL y RANKING_PORTEROS son valores: se rehacen desde las hojas (misma lógica que ranking.py)."""
    def tot(v, cols):
        return sum(v[c - 1] if isinstance(v[c - 1], (int, float)) else 0 for c in cols)
    field, gks = [], []
    for sh in SHEETS:
        ws = wb[sh]
        for r in range(4, ws.max_row + 1):
            v = [ws.cell(r, c).value for c in range(1, 37)]
            if not v[0] or not v[4]:
                continue
            T = dict(pj=tot(v, (7, 11, 16, 20)), mn=tot(v, (8, 12, 17, 21)), g=tot(v, (9, 13, 18, 22)), a=tot(v, (10, 14, 19, 23)))
            (gks if sh == "PORTEROS" else field).append((v, T, sh))
    styles = {}
    for sh in SHEETS:
        ws = wb[sh]
        for r in range(4, ws.max_row + 1):
            for c in (5, 6):
                x = ws.cell(r, c)
                if x.value and (c, x.value) not in styles:
                    styles[(c, x.value)] = copy.copy(x._style)
    ws = wb["RANKING_TOTAL"]
    hdr = [c.value for c in ws[3]]
    ncol = len([h for h in hdr if h])
    tmpl = [copy.copy(ws.cell(4, c)._style) for c in range(1, ncol + 1)]
    field.sort(key=lambda x: (-(x[1]["g"] + x[1]["a"]), -x[1]["g"], -x[1]["mn"]))
    ws.delete_rows(4, ws.max_row - 3)
    origin = "Puesto origen" in hdr
    for i, (v, T, sh) in enumerate(field):
        row = [i + 1, v[0], v[1], v[4], v[5]] + ([sh] if origin else []) + [
            T["pj"], T["mn"], T["g"], T["a"], T["g"] + T["a"], round(T["g"] / T["mn"] * 90, 2) if T["mn"] else 0,
            round(T["mn"] / T["pj"], 1) if T["pj"] else 0, v[34]]
        for c, val in enumerate(row, 1):
            cell = ws.cell(4 + i, c)
            cell.value, cell._style = val, copy.copy(tmpl[c - 1])
        for c, src in ((4, 5), (5, 6)):
            if (src, v[src - 1]) in styles:
                ws.cell(4 + i, c)._style = copy.copy(styles[(src, v[src - 1])])
    _refresh_cf(ws, 3 + len(field), ncol)
    wp = wb["RANKING_PORTEROS"]
    tm = [copy.copy(wp.cell(4, c)._style) for c in range(1, 13)]
    gks.sort(key=lambda x: (-x[1]["a"], x[1]["g"] if x[1]["pj"] else 999, -x[1]["pj"]))
    wp.delete_rows(4, wp.max_row - 3)
    for i, (v, T, _) in enumerate(gks):
        row = [i + 1, v[0], v[4], v[5], T["pj"], T["mn"], T["g"], T["a"], (T["a"] / T["pj"]) if T["pj"] else 0,
               (T["g"] / T["mn"] * 90) if T["mn"] else 0, (T["mn"] / T["pj"]) if T["pj"] else 0, v[34]]
        for c, val in enumerate(row, 1):
            cell = wp.cell(4 + i, c)
            cell.value, cell._style = val, copy.copy(tm[c - 1])
        for c, src in ((3, 5), (4, 6)):
            if (src, v[src - 1]) in styles:
                wp.cell(4 + i, c)._style = copy.copy(styles[(src, v[src - 1])])
    _refresh_cf(wp, 3 + len(gks), 12)


def _refresh_cf(ws, last, ncol):
    import re
    old = [(str(c.sqref), c.rules) for c in ws.conditional_formatting]
    ws.conditional_formatting = type(ws.conditional_formatting)()
    for sq, rules in old:
        c0 = re.match(r"([A-Z]+)", sq).group(1)
        for rule in rules:
            ws.conditional_formatting.add(f"{c0}4:{c0}{last}", rule)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(ncol)}{last}"


if __name__ == "__main__":
    main()
