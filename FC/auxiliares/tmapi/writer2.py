import openpyxl, json, copy, collections, re, pickle
F=json.load(open('final.json'))
SH=['DELANTEROS','EXTREMOS','MEDIAPUNTAS','MEDIOCENTROS','DEFENSAS','PORTEROS']
FILES=[('Temporada 2025-26.xlsx','2526'),('Temporada 2026-27.xlsx','2627')]
# ---- TM player info
s=open('pi.txt').read().rstrip('#'); a,b=s.split('~~')
P={x.split('|')[0]:x.split('|') for x in a.split('~')}
K={x.split('|')[1]:x.split('|') for x in b.split('~')}
for x in open('meta.txt').read().split('~'):
    f=x.split('|')
    if f[0]=='K' and f[3]=='1': K.setdefault(f[1],f)
C2A={k[5]:k[6] for k in K.values() if k[1]==k[4] and k[6]}
POSMAP={'CF':'DC','SS':'MCO','LW':'EI','RW':'ED','LM':'EI','RM':'ED','AM':'MCO','CM':'MC','DM':'MCD','CB':'DFC','LB':'LI','RB':'LD','GK':'POR'}
SHEETOF={'DC':'DELANTEROS','EI':'EXTREMOS','ED':'EXTREMOS','MCO':'MEDIAPUNTAS','MC':'MEDIOCENTROS','MCD':'MEDIOCENTROS','DFC':'DEFENSAS','LI':'DEFENSAS','LD':'DEFENSAS','POR':'PORTEROS'}
def nat(pid,cur=None):
    p=P.get(pid)
    if not p: return None
    nt=p[3]
    c=None
    if nt and nt in K and K[nt][4]==nt: c=K[nt][6]          # selección absoluta
    if not c:
        opts=[C2A.get(p[1]),C2A.get(p[2])]
        if cur and cur in opts: return cur                  # sin absoluta: se respeta si es una de sus nacionalidades
        c=opts[0]
    if not c and nt and nt in K: c=(K.get(K[nt][4]) or [None]*7)[6]
    return c.upper() if c else None
def mval(pid,tag):
    p=P.get(pid)
    if not p: return None
    v=p[6] if tag=='2627' else (p[11] if len(p)>11 else '')
    if not v: return None
    x=int(v)/1e6
    return round(x,3) if x<1 else round(x,2)
AGE={}
for ent in open('sq.txt').read().split('~'):
    for q in ent.split('#')[1:]:
        f=q.split('|')
        if len(f)>3 and f[3].isdigit(): AGE.setdefault(f[0],int(f[3]))
wbs={f:openpyxl.load_workbook(f) for f,_ in FILES}
log=collections.Counter(); changes=collections.defaultdict(list)
out_rows={}
for fname,tag in FILES:
    wb=wbs[fname]
    clubst={}; ligst={}; contst={}; tmpl={}
    for sh in SH:
        ws=wb[sh]; tmpl[sh]=[copy.copy(ws.cell(4,c)._style) for c in range(1,37)]
        for r in range(4,ws.max_row+1):
            e=ws.cell(r,5); f=ws.cell(r,6); o=ws.cell(r,15)
            if e.value and e.value not in clubst: clubst[e.value]=copy.copy(e._style)
            if f.value and f.value not in ligst: ligst[f.value]=copy.copy(f._style)
            if o.value and o.value not in contst: contst[o.value]=copy.copy(o._style)
    items={sh:[] for sh in SH}; notes={sh:[] for sh in SH}; maxr={}
    for sh in SH:
        ws=wb[sh]; maxr[sh]=ws.max_row
        for r in range(4,ws.max_row+1):
            vals=[ws.cell(r,c).value for c in range(1,37)]
            if all(v is None for v in vals): continue
            sty=[copy.copy(ws.cell(r,c)._style) for c in range(1,37)]
            k=f"{tag}|{sh}|{r}"; rec=F.get(k)
            it=dict(vals=vals,sty=sty,rec=rec,sh=sh)
            if rec is None or (rec['status']=='unmatched' and vals[4] is None): it['rec']=None; notes[sh].append(it); continue
            if rec['status']=='dupe': log['dupe_removed']+=1; continue
            variants=[(vals,sty,rec)]+[(copy.deepcopy(vals),[copy.copy(z) for z in sty],e) for e in rec.get('extra',[])]
            for vals,sty,rec in variants:
                target=sh
                if rec['status']=='unmatched':
                    for c in list(range(7,15))+list(range(16,28)): vals[c-1]=0
                    log['unmatched_zero']+=1
                else:
                    pid=rec['pid']; pinfo=P.get(pid)
                    # --- position
                    if pinfo and pinfo[5] in POSMAP:
                        newpos=POSMAP[pinfo[5]]
                        if newpos!=vals[1]: changes['pos'].append((tag,vals[0],vals[1],newpos)); vals[1]=newpos
                        target=SHEETOF[newpos]
                    # --- nationality
                    n=nat(pid,vals[3])
                    if n and n!=vals[3]: changes['nac'].append((tag,vals[0],vals[3],n)); vals[3]=n
                    # --- value
                    mv=mval(pid,tag)
                    if mv is not None and mv!=vals[34]: changes['valor'].append((tag,vals[0],vals[34],mv)); vals[34]=mv
                    # --- age
                    if vals[2] in (None,'') and pid in AGE: vals[2]=AGE[pid]
                    # --- stats
                    o=rec['out']; gk=(target=='PORTEROS'); idx=(0,1,4,5) if gk else (0,1,2,3)
                    for bk,start in (('LIGA',7),('COPA',11),('CONT',16),('FIFA',20),('SEL',24)):
                        for j,i in enumerate(idx): vals[start-1+j]=o[bk][i]
                    if rec.get('cont'):
                        vals[14]=rec['cont']
                        if rec['cont'] in contst: sty[14]=copy.copy(contst[rec['cont']])
                    if rec.get('club'):
                        vals[4]=rec['club']; sty[4]=copy.copy(clubst.get(rec['club'],sty[3]))
                        if rec.get('liga'): vals[5]=rec['liga']; sty[5]=copy.copy(ligst.get(rec['liga'],sty[3]))
                        log['club_changed']+=1
                    log['updated']+=1
                if rec.get('conf') and rec is not variants[0][2]: log['extra_continent']+=1
                if target!=sh:
                    t=tmpl[target]
                    sty=[copy.copy(t[c]) if c not in (4,5,14,35) else sty[c] for c in range(36)]
                    log['moved_sheet']+=1; changes['moved'].append((tag,vals[0],sh,target))
                it=dict(vals=vals,sty=sty,rec=rec,sh=target)
                items[target].append(it)
    def tot(v,cols): return sum(v[i-1] if isinstance(v[i-1],(int,float)) else 0 for i in cols)
    allrows={}
    for sh in SH:
        ws=wb[sh]; gk= sh=='PORTEROS'; data=items[sh]
        for it in data:
            v=it['vals']; it['T']=dict(pj=tot(v,(7,11,16,20)),mn=tot(v,(8,12,17,21)),g=tot(v,(9,13,18,22)),a=tot(v,(10,14,19,23)))
        if gk: data.sort(key=lambda it:(-it['T']['a'],it['T']['g'] if it['T']['pj'] else 999,-it['T']['pj'],-it['T']['mn']))
        else: data.sort(key=lambda it:(-(it['T']['g']+it['T']['a']),-it['T']['g'],-it['T']['mn']))
        allrows[sh]=data
        out=data+notes[sh]
        mr=maxr[sh]
        for r in range(4,max(mr,3+len(out))+1):
            for c in range(1,37): ws.cell(r,c).value=None
        for i,it in enumerate(out):
            r=4+i; v=it['vals']
            for c in range(1,37):
                cell=ws.cell(r,c); cell._style=it['sty'][c-1]
                if it['rec'] is None:
                    cell.value = v[0] if c==1 else None
                elif not (28<=c<=34): cell.value=v[c-1]
            if it['rec'] is None: continue
            ws[f'AB{r}']=f'=G{r}+K{r}+P{r}+T{r}'; ws[f'AC{r}']=f'=H{r}+L{r}+Q{r}+U{r}'
            ws[f'AD{r}']=f'=I{r}+M{r}+R{r}+V{r}'; ws[f'AE{r}']=f'=J{r}+N{r}+S{r}+W{r}'
            ws[f'AF{r}']=f'=IF(AB{r}=0,0,AE{r}/AB{r})' if gk else f'=AD{r}+AE{r}'
            ws[f'AG{r}']=f'=IF(AC{r}=0,0,AD{r}/AC{r}*90)'; ws[f'AH{r}']=f'=IF(AB{r}=0,0,AC{r}/AB{r})'
        last=3+len(out)
        if ws.max_row>last: ws.delete_rows(last+1,ws.max_row-last)
        old=[(str(cf.sqref),cf.rules) for cf in ws.conditional_formatting]
        ws.conditional_formatting=type(ws.conditional_formatting)()
        for sq,rules in old:
            c0=re.match(r'([A-Z]+)',sq).group(1)
            for rule in rules: ws.conditional_formatting.add(f'{c0}4:{c0}{3+len(data)}',rule)
        ws.auto_filter.ref=f'A3:AJ{last}'
    out_rows[fname]={sh:[dict(vals=it['vals'],T=it['T']) for it in allrows[sh]] for sh in SH}
    wb.save('stage1_'+fname)
    print(fname, dict(log)); log.clear()
pickle.dump(out_rows,open('allrows.pkl','wb'))
json.dump({k:v for k,v in changes.items()},open('changes.json','w'),ensure_ascii=False)
for k,v in changes.items(): print(k,len(v),v[:12])
