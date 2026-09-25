import openpyxl, copy, pickle, re
SH=['DELANTEROS','EXTREMOS','MEDIAPUNTAS','MEDIOCENTROS','DEFENSAS','PORTEROS']
AR=pickle.load(open('allrows.pkl','rb'))
for fname in ['Temporada 2025-26.xlsx','Temporada 2026-27.xlsx']:
    wb=openpyxl.load_workbook('stage1_'+fname)
    clubst={}; ligst={}
    for sh in SH:
        ws=wb[sh]
        for r in range(4,ws.max_row+1):
            e=ws.cell(r,5); f=ws.cell(r,6)
            if e.value and e.value not in clubst: clubst[e.value]=copy.copy(e._style)
            if f.value and f.value not in ligst: ligst[f.value]=copy.copy(f._style)
    # ---- RANKING_TOTAL
    ws=wb['RANKING_TOTAL']; hdr=[c.value for c in ws[3]]; ncol=len([h for h in hdr if h])
    tmpl=[copy.copy(ws.cell(4,c)._style) for c in range(1,ncol+1)]
    field=[]
    for sh in SH[:-1]:
        for it in AR[fname][sh]:
            v=it['vals']; T=it['T']
            field.append((v,T,sh))
    field.sort(key=lambda x:(-(x[1]['g']+x[1]['a']),-x[1]['g'],-x[1]['mn']))
    maxr=ws.max_row
    ws.delete_rows(4,maxr-3)
    has_origin = 'Puesto origen' in hdr
    for i,(v,T,sh) in enumerate(field):
        r=4+i; g90=round(T['g']/T['mn']*90,2) if T['mn'] else 0; mpj=round(T['mn']/T['pj'],1) if T['pj'] else 0
        row=[i+1,v[0],v[1],v[4],v[5]]+([sh] if has_origin else [])+[T['pj'],T['mn'],T['g'],T['a'],T['g']+T['a'],g90,mpj,v[34]]
        for c,val in enumerate(row,1):
            cell=ws.cell(r,c); cell.value=val; cell._style=copy.copy(tmpl[c-1])
        if v[4] in clubst: ws.cell(r,4)._style=copy.copy(clubst[v[4]])
        if v[5] in ligst: ws.cell(r,5)._style=copy.copy(ligst[v[5]])
    last=3+len(field)
    old=[(str(cf.sqref),cf.rules) for cf in ws.conditional_formatting]
    ws.conditional_formatting=type(ws.conditional_formatting)()
    for sq,rules in old:
        c0=re.match(r'([A-Z]+)',sq).group(1)
        for rule in rules: ws.conditional_formatting.add(f'{c0}4:{c0}{last}',rule)
    ws.auto_filter.ref=f'A3:{openpyxl.utils.get_column_letter(ncol)}{last}'
    # ---- RANKING_PORTEROS
    wp=wb['RANKING_PORTEROS']; tm=[copy.copy(wp.cell(4,c)._style) for c in range(1,13)]
    gks=[(it['vals'],it['T']) for it in AR[fname]['PORTEROS']]
    gks.sort(key=lambda x:(-x[1]['a'], x[1]['g'] if x[1]['pj'] else 999, -x[1]['pj']))
    wp.delete_rows(4,wp.max_row-3)
    for i,(v,T) in enumerate(gks):
        r=4+i
        row=[i+1,v[0],v[4],v[5],T['pj'],T['mn'],T['g'],T['a'],(T['a']/T['pj']) if T['pj'] else 0,(T['g']/T['mn']*90) if T['mn'] else 0,(T['mn']/T['pj']) if T['pj'] else 0,v[34]]
        for c,val in enumerate(row,1):
            cell=wp.cell(r,c); cell.value=val; cell._style=copy.copy(tm[c-1])
        if v[4] in clubst: wp.cell(r,3)._style=copy.copy(clubst[v[4]])
        if v[5] in ligst: wp.cell(r,4)._style=copy.copy(ligst[v[5]])
    lp=3+len(gks)
    old=[(str(cf.sqref),cf.rules) for cf in wp.conditional_formatting]
    wp.conditional_formatting=type(wp.conditional_formatting)()
    for sq,rules in old:
        c0=re.match(r'([A-Z]+)',sq).group(1)
        for rule in rules: wp.conditional_formatting.add(f'{c0}4:{c0}{lp}',rule)
    if wp.auto_filter.ref: wp.auto_filter.ref=f'A3:L{lp}'
    wb.save('stage2_'+fname)
    print(fname,'ranking',len(field),'porteros',len(gks))
