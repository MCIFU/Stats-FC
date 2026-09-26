"""Completa ligas en Temporada 2025-26.xlsx y Temporada 2026-27.xlsx con TODOS los jugadores que han jugado minutos de
liga (Transfermarkt partido a partido, mismas reglas que actualizar_excels.py).

Uso (desde la carpeta FC):
    python auxiliares/tmapi/ampliar_ligas.py --ligas "Süper Lig,Liga Belga"        # informe, no escribe
    python auxiliares/tmapi/ampliar_ligas.py --ligas "Süper Lig,Liga Belga" --escribir

- Candidatos: plantillas TM de todos los clubes de la liga en esa temporada. Entra quien tenga partidos de LIGA con un
  club de esa liga y NO tenga ya fila en el Excel de esa temporada en el mismo continente (regla de continentes).
- Club = club de la liga en su último partido de liga de la temporada. Liga = etiqueta del Excel.
- Datos de la ficha TM: nombre (el de la plantilla TM, igual que el cruce), edad, posición (hoja), NAC (selección absoluta
  o 1.ª nacionalidad, código FIFA), valor de mercado (25-26: el vigente a 31/07/2026) y contrato (solo 26-27).
- Estadísticas LIGA/COPA/CONT/FIFA/SEL calculadas igual que el resto de filas. Después: estilo_temporadas.py ordena y da formato.
"""
import argparse, collections, copy, datetime as dt, json, re, sys, time
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import descargar_partidos as D  # noqa: E402
import actualizar_excels as A  # noqa: E402

FC = D.FC
# etiqueta de liga en el Excel -> (competición TM para la lista de clubes, año natural)
LIGAS = {"Eredivisie": ("NL1", False), "Süper Lig": ("TR1", False), "Liga Belga": ("BE1", False), "Saudi Pro": ("SA1", False),
         "Liga MX": ("MEXA", False), "Scottish Premiership": ("SC1", False), "Liga Argentina": ("ARG1", True),
         "Brasileirão": ("BRA1", True), "Ekstraklasa": ("PL1", False), "MLS": ("MLS1", True), "Portugal": ("PO1", False),
         "Championship": ("GB2", False), "LaLiga2": ("ES2", False), "Serie B": ("IT2", False), "2. Bundesliga": ("L2", False),
         "Super League 1": ("GR1", False), "Superliga": ("DK1", False), "Super League": ("C1", False),
         "Bundesliga Austria": ("A1", False), "J1 League": ("JAP1", True), "K League 1": ("RSK1", True)}
# Japón pasa a temporada ago-may en 2026/27 (TM: saison_id 2026); el torneo corto feb-jun 2026 es otra competición
SID = {("JAP1", "2026-27"): 2026}
POSMAP = {"CF": "DC", "SS": "MCO", "LW": "EI", "RW": "ED", "LM": "EI", "RM": "ED", "AM": "MCO", "CM": "MC", "DM": "MCD",
          "CB": "DFC", "LB": "LI", "RB": "LD", "GK": "POR"}
GROUPMAP = {"GOALKEEPER": "POR", "DEFENDER": "DFC", "MIDFIELD": "MC", "FORWARD": "DC"}
SHEETOF = {"DC": "DELANTEROS", "EI": "EXTREMOS", "ED": "EXTREMOS", "MCO": "MEDIAPUNTAS", "MC": "MEDIOCENTROS",
           "MCD": "MEDIOCENTROS", "DFC": "DEFENSAS", "LI": "DEFENSAS", "LD": "DEFENSAS", "POR": "PORTEROS"}


def hexcol(c, default):
    """Color de TM a RRGGBB ('#fff' → 'FFFFFF'); lo que no sea hexadecimal vale el color por defecto."""
    c = (c or "").strip().lstrip("#")
    if re.fullmatch(r"[0-9A-Fa-f]{3}", c):
        c = "".join(x * 2 for x in c)
    return c[:6].upper() if re.fullmatch(r"[0-9A-Fa-f]{6}([0-9A-Fa-f]{2})?", c) else default


def clubes_liga(comp, sid):
    def fetch():
        h = D.get(f"{D.TMWEB}/x/startseite/wettbewerb/{comp}/saison_id/{sid}", as_json=False)
        return list(dict.fromkeys(re.findall(r'href="/[^"/]+/startseite/verein/(\d+)/saison_id/' + str(sid), h)))
    return D.cached(f"ligas/{comp}_{sid}.json.gz", fetch)


def fichas(pids):
    """tmapi /players en lotes de 50 (caché acumulativa)."""
    import gzip
    p = D.CACHE / "meta_players.json.gz"
    have = json.loads(gzip.decompress(p.read_bytes())) if p.exists() else {}
    need = sorted(set(pids) - set(have))
    for i in range(0, len(need), 50):
        q = "&".join(f"ids%5B%5D={x}" for x in need[i:i + 50])
        for it in D.as_list(D.get(f"{D.TMAPI}/players?{q}")):
            have[str(it.get("id"))] = it
        time.sleep(0.3)
    p.write_bytes(gzip.compress(json.dumps(have, ensure_ascii=False).encode()))
    return have


def valor(f, season):
    mv = f.get("marketValueDetails") or {}
    cur, prev = mv.get("current") or {}, mv.get("previous") or {}
    pick = cur
    if season == "2025-26" and str(cur.get("determined") or "") > "2026-07-31" and prev.get("value") is not None:
        pick = prev
    v = pick.get("value")
    return None if v is None else (round(v / 1e6, 3) if v < 1e6 else round(v / 1e6, 2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ligas", required=True, help="etiquetas del Excel separadas por comas")
    ap.add_argument("--corte", default="2026-09-24")
    ap.add_argument("--escribir", action="store_true")
    ap.add_argument("--workers", type=int, default=6)
    a = ap.parse_args()
    ligas = [l.strip() for l in a.ligas.split(",")]
    for l in ligas:
        if l not in LIGAS:
            sys.exit(f"Liga no configurada: {l}. Opciones: {', '.join(LIGAS)}")

    # 1) filas actuales (para no duplicar) y candidatos por liga y temporada
    rows = A.excel_rows()
    rows = D.map_ids([r for r in rows if r["club_id"]])
    exist = collections.defaultdict(list)  # (temporada, tm_id) -> club_ids de sus filas
    for r in rows:
        if r.get("tm_id"):
            exist[(r["season"], r["tm_id"])].append(r["club_id"])
    cand = {}  # (temporada, liga) -> {tm_id: nombre plantilla}
    league_clubs = {}
    for liga in ligas:
        comp, cal = LIGAS[liga]
        for season in ("2025-26", "2026-27"):
            sid = (2024 if season == "2025-26" else 2025) if cal else (2025 if season == "2025-26" else 2026)
            sid = SID.get((comp, season), sid)
            clubs = clubes_liga(comp, sid)
            league_clubs[(season, liga)] = set(clubs)
            names = {}
            for c in clubs:
                fresh = not (D.CACHE / f"squads/{c}_{sid}.json.gz").exists()
                names.update(D.squad(c, sid))
                if fresh:
                    time.sleep(1.5)
            cand[(season, liga)] = names
            print(f"{season} {liga}: {len(clubs)} clubes, {len(names)} jugadores en plantillas")
    pids = sorted({r["tm_id"] for r in rows if r.get("tm_id")} | {p for d in cand.values() for p in d})
    all_clubs = {r["club_id"] for r in rows} | {c for s in league_clubs.values() for c in s}
    ctx = A.construir(pids, False, a.workers, a.corte, all_clubs)
    CLUB = ctx["CLUB"]

    # 2) quién entra: partidos de LIGA con un club de la liga en esa temporada y sin fila del mismo continente
    nuevos = []
    for (season, liga), names in cand.items():
        lc = league_clubs[(season, liga)]
        if not lc:
            print(f"  {season} {liga}: sin clubes en TM, se omite")
            continue
        conf = collections.Counter(CLUB.get(c, {}).get("conf") for c in lc).most_common(1)[0][0]
        for pid, nm in names.items():
            G = [x for x in ctx["valid"].get(pid, []) if x["tag"] == season and x["blk"] == "LIGA" and x["club"] in lc]
            if not G:
                continue
            if any(CLUB.get(c, {}).get("conf") == conf for c in exist.get((season, pid), [])):
                continue  # ya tiene fila en este continente (p.ej. otra liga del Excel)
            last = max(G, key=lambda x: x["date"])["club"]
            nuevos.append(dict(season=season, liga=liga, tm_id=pid, name=nm, club_id=last))
            exist[(season, pid)].append(last)
    print(f"jugadores a añadir: {len(nuevos)}", dict(collections.Counter((n['season'], n['liga']) for n in nuevos)))
    if not nuevos:
        return

    # 3) ficha TM y estadísticas
    F = fichas({n["tm_id"] for n in nuevos})
    nt_abbr = {}
    for k, v in ctx["KM"].items():
        b = v.get("baseDetails") or {}
        if b.get("isNationalTeam") and str(b.get("mainClubId") or k) == k and b.get("abbreviation"):
            nt_abbr[b.get("countryId")] = b["abbreviation"]
    nt_ids = {ca["clubId"] for f in F.values() for ca in (f.get("clubAssignments") or []) if ca.get("type") == "nationalTeam"}
    KM2 = D.meta("clubs", nt_ids)
    rev = json.loads((HERE / "clubmap.json").read_text(encoding="utf-8"))
    excel_name = {}
    for n_, i_ in rev.items():
        excel_name.setdefault(str(i_), n_)
    for n in nuevos:
        f = F.get(n["tm_id"], {})
        at = f.get("attributes") or {}
        pos = POSMAP.get(((at.get("position") or {}).get("shortName") or "").upper()) or GROUPMAP.get(at.get("positionGroup"), "MC")
        nat = None
        for ca in f.get("clubAssignments") or []:
            if ca.get("type") == "nationalTeam":
                k = (KM2.get(str(ca["clubId"])) or {}).get("baseDetails") or {}
                nat = (KM2.get(str(k.get("mainClubId"))) or {}).get("baseDetails", {}).get("abbreviation") or k.get("abbreviation")
        if not nat:
            nid = ((f.get("nationalityDetails") or {}).get("nationalities") or {}).get("nationalityId")
            nat = nt_abbr.get(nid)
        n.update(pos=pos, sheet=SHEETOF[pos], gk=pos == "POR", age=(f.get("lifeDates") or {}).get("age"), nat=nat,
                 value=valor(f, n["season"]), contract=int(str(at.get("contractUntil"))[:4]) if at.get("contractUntil") else None,
                 club=excel_name.get(n["club_id"]) or CLUB.get(n["club_id"], {}).get("name") or n["club_id"])
    # guardar nombre -> ID de los clubes nuevos (el nombre corto TM puede repetirse en otro país: Al-Ittihad)
    ce = HERE / "clubmap_extra.json"
    ext = json.loads(ce.read_text(encoding="utf-8")) if ce.exists() else {}
    for n in nuevos:
        if n["club"] not in rev and n["club"] not in ext:
            ext[n["club"]] = n["club_id"]
    ce.write_text(json.dumps(ext, ensure_ascii=False, indent=0), encoding="utf-8")
    res, _ = A.agregar([dict(n, row=None, old=None) for n in nuevos], ctx)
    stats = {(r["season"], r["tm_id"], r["excel_club"]): r["new"] for r in res}  # un jugador puede tener 2 filas (2 continentes)

    if not a.escribir:
        for n in nuevos[:15]:
            print(" ", n["season"], n["liga"], n["name"], n["pos"], n["age"], n["nat"], n["club"], n["value"], stats[(n["season"], n["tm_id"], n["club"])][:4])
        return

    # 4) escribir al final de su hoja (estilo_temporadas.py ordena y da formato después)
    for season, fname in A.FILES.items():
        wb = openpyxl.load_workbook(FC / fname)
        styles = {}
        for sh in A.SHEETS:
            ws = wb[sh]
            for r in range(4, ws.max_row + 1):
                for c in (5, 6):
                    v = ws.cell(r, c).value
                    if v and (c, v) not in styles:
                        styles[(c, v)] = (copy.copy(ws.cell(r, c).fill), copy.copy(ws.cell(r, c).font))
        has_contract = wb["DELANTEROS"].cell(3, 36).value == "Contrato"
        count = collections.Counter()
        for n in [x for x in nuevos if x["season"] == season]:
            ws = wb[n["sheet"]]
            r = ws.max_row + 1
            while r > 4 and not ws.cell(r - 1, 1).value:
                r -= 1
            if ws.cell(r, 1).value:  # hay una nota al pie: insertar antes
                ws.insert_rows(r)
            new = stats[(season, n["tm_id"], n["club"])]
            vals = [n["name"], n["pos"], n["age"], n["nat"], n["club"], n["liga"]] + [x if x is not None else ("—" if i == 8 else 0) for i, x in enumerate(new)]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c).value = v
            ws.cell(r, 35).value = n["value"]
            if has_contract:
                ws.cell(r, 36).value = n["contract"]
                if n["contract"] in (2027, 2028):
                    ws.cell(r, 36).fill = PatternFill("solid", fgColor="DC2626" if n["contract"] == 2027 else "EA580C")
                    ws.cell(r, 36).font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            for c, key in ((5, n["club"]), (6, n["liga"])):
                if (c, key) in styles:
                    ws.cell(r, c).fill, ws.cell(r, c).font = styles[(c, key)]
                elif c == 5:  # club nuevo: sus colores de Transfermarkt
                    cols = (((ctx["KM"].get(n["club_id"]) or {}).get("baseDetails") or {}).get("superiorClub") or {}).get("colors") or {}
                    bg, fg = hexcol(cols.get("firstColor"), "334155"), hexcol(cols.get("secondColor"), "FFFFFF")
                    if bg.upper() == fg.upper():
                        fg = "FFFFFF" if bg.upper() != "FFFFFF" else "0F172A"
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=bg)
                    ws.cell(r, c).font = Font(name="Arial", size=9, bold=True, color=fg)
                    styles[(5, key)] = (copy.copy(ws.cell(r, c).fill), copy.copy(ws.cell(r, c).font))
            count[n["liga"]] += 1
        A.rebuild_rankings(wb)
        L = [s for s in wb.sheetnames if s.startswith("LEEME")][0]
        wsl = wb[L]
        c = wsl.cell(wsl.max_row + 1, 1)
        c.value = (f"AMPLIACIÓN {dt.date.today():%d/%m/%Y} (Claude): añadidos todos los jugadores con minutos de liga según Transfermarkt "
                   f"(corte {a.corte}): " + ", ".join(f"{k} +{v}" for k, v in count.most_common()) +
                   ". Edad/posición/NAC/valor/contrato de la ficha TM; estadísticas partido a partido con las reglas de siempre.")
        wb.save(FC / fname)
        print(f"{fname}: {sum(count.values())} filas añadidas", dict(count))


if __name__ == "__main__":
    main()
