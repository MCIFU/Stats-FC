"""Genera el Excel de valoración a partir del resultado del motor.
- Capas de datos (percentiles, atributos, roles...) se escriben como VALORES (se regeneran con el script).
- La capa final (CA, PA, SCOUT) se escribe con FÓRMULAS que leen los parámetros de la hoja PARAMETROS:
  si cambias un peso allí, Excel recalcula CA/PA/SCOUT al instante.
"""
import json
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule, ColorScaleRule
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

import engine

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_GREY = Font(name=FONT, size=9, italic=True, color="666666")
FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_UNK = PatternFill("solid", fgColor="EDEDED")
FILL_FORM = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")


def hdr(ws, row, cols, start=1, fill=FILL_HDR, font=F_HDR):
    for j, c in enumerate(cols, start):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill = font, fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_df(ws, df, start_row=1, num_fmt=None, widths=None, freeze="B2", autofilter=True):
    hdr(ws, start_row, list(df.columns))
    vals = df.astype(object).where(df.notna(), None).values.tolist()
    for i, row in enumerate(vals, start_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating,)):
                v = float(v)
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            ws.cell(row=i, column=j, value=v)
    if num_fmt:
        for j, col in enumerate(df.columns, 1):
            f = num_fmt(col)
            if f:
                for i in range(start_row + 1, start_row + 1 + len(df)):
                    ws.cell(row=i, column=j).number_format = f
    for j, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col, max(9, min(28, len(str(col)) + 2)))
        ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = freeze
    if autofilter and len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    ws.row_dimensions[start_row].height = 42


def grey_blanks(ws, ref):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEN({ref.split(":")[0].replace("$", "")})=0'], fill=FILL_UNK))


def scale(ws, ref):
    ws.conditional_formatting.add(ref, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                      mid_type="num", mid_value=50, mid_color="FFEB84",
                                                      end_type="num", end_value=100, end_color="63BE7B"))


# ------------------------------------------------------------------ PARAMETROS (entradas editables)
def params_sheet(wb, cfg):
    ws = wb.create_sheet("PARAMETROS")
    M = cfg["model"]
    ws["A1"] = "PARÁMETROS DEL MODELO (celdas amarillas = editables; CA/PA/SCOUT se recalculan solos)"
    ws["A1"].font = F_TITLE
    ws["A2"] = f"Versión: {M['model_version']}  ·  Si cambias algo aquí, cambia también model_version para no mezclar valoraciones."
    ws["A2"].font = F_GREY
    rows = [
        ("model_version", M["model_version"], "Versión del modelo. Cada valoración guarda la suya."),
        ("k_min_confianza", M["sample"]["confidence_k_minutes"], "Confianza de muestra = 1 - EXP(-Min/k)"),
        ("min_minutos_valorar", M["sample"]["min_minutes_rated"], "Por debajo: sin CA (UNKNOWN)"),
        ("factor_base_liga", M["current_ability"]["league_base_factor"], "Nivel del jugador mediano de una liga = factor × CONTEXT_SCORE"),
        ("pendiente_contexto", M["current_ability"]["context_slope"], "CA_CONTEXT = base + pendiente × (score en su liga − 50)"),
        ("centro_raw", M["current_ability"]["raw_scale_center"], "CA_RAW = centro + pendiente_raw × (score global − 50)"),
        ("pendiente_raw", M["current_ability"]["raw_scale_slope"], ""),
        ("peso_contexto", M["current_ability"]["w_context"], "Mezcla CA_CONTEXT vs CA_RAW"),
        ("peso_raw", M["current_ability"]["w_raw"], ""),
        ("conf_peso_muestra", M["confidence"]["w_sample"], "CA_CONFIDENCE = muestra + cobertura de datos + calidad de datos de la liga"),
        ("conf_peso_cobertura", M["confidence"]["w_coverage"], ""),
        ("conf_peso_calidad_liga", M["confidence"]["w_competition_data"], ""),
        ("pa_mult_min", M["potential"]["relperf_multiplier_min"], "Crecimiento efectivo = crecimiento_edad × (min + (max−min)×REL_PERF/100) × margen"),
        ("pa_mult_max", M["potential"]["relperf_multiplier_max"], ""),
        ("pa_margen_ref", M["potential"]["headroom_ref"], "margen = MIN(1, (100−CA)/margen_ref)"),
        ("pa_peso_tendencia", M["potential"]["trend_weight"], "Suma peso × (CA 26-27 − CA 25-26), acotado ±5"),
    ]
    for k, v in M["scout_score"]["weights"].items():
        rows.append((f"scout_{k}", v, "Peso en SCOUT SCORE (si FORM es UNKNOWN se reparte su peso)"))
    hdr(ws, 4, ["Parámetro", "Valor", "Explicación"])
    names = {}
    for i, (k, v, e) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=k).font = F_BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.fill = F_INPUT, FILL_INPUT
        ws.cell(row=i, column=3, value=e).font = F_GREY
        names[k] = f"PARAMETROS!$B${i}"
    # tablas por edad (enteros) — mismas que usa el motor
    P = M["potential"]
    ages = list(range(15, 46))
    r0 = 5
    ws.cell(row=r0 - 1, column=5, value="Edad")
    hdr(ws, r0 - 1, ["Edad", "Crecimiento máx.", "Ancho rango PA", "Confianza por edad", "Puntuación edad (scout)"], start=5)
    for i, a in enumerate(ages, r0):
        ws.cell(row=i, column=5, value=a).font = F_BOLD
        for j, t in enumerate([P["growth_by_age"], P["range_base"], P["age_confidence"], M["scout_score"]["age_score"]], 6):
            c = ws.cell(row=i, column=j, value=round(float(engine.interp_age(t, a)), 3))
            c.font, c.fill = F_INPUT, FILL_INPUT
    last = r0 + len(ages) - 1
    names["tabla_edad"] = f"PARAMETROS!$E${r0}:$I${last}"
    for col, w in zip("ABCDEFGHI", [26, 18, 80, 3, 8, 16, 16, 18, 22]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    return names


# ------------------------------------------------------------------ hojas principales
ROLE_LABELS = None


def build(base, cfg, out_path, raw=None, sims=None, corr=None, matches=None, history=None, source_note=""):
    global ROLE_LABELS
    M = cfg["model"]
    attrs = cfg["attributes"]["attributes"]
    roles = cfg["roles"]["roles"]
    metrics = cfg["metrics"]["metrics"]
    ROLE_LABELS = {r: d["label"] for r, d in roles.items()}
    b = base.copy()
    b = b.sort_values(["season", "pos_group", "CA_FINAL"], ascending=[True, True, False], na_position="last").reset_index(drop=True)
    b["KEY"] = b.name.astype(str) + " | " + b.club.astype(str) + " | " + b.season.astype(str)
    # rol principal
    rcols = [f"role_{r}" for r in roles]
    rv = b[rcols]
    has = rv.notna().any(axis=1)
    b["best_role"] = None
    if has.any():
        b.loc[has, "best_role"] = rv[has].idxmax(axis=1).map(lambda x: ROLE_LABELS.get(str(x)[5:]))
    b["best_role_score"] = rv.max(axis=1)

    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    wb.remove(wb.active)
    leeme = wb.create_sheet("LEEME")
    ficha = wb.create_sheet("FICHA")
    comp = wb.create_sheet("COMPARADOR")
    P = params_sheet(wb, cfg)

    # ---------------- VALORACION (fórmulas en la capa final)
    ws = wb.create_sheet("VALORACION")
    cols = [("KEY", "KEY"), ("ID", "player_id"), ("Jugador", "name"), ("Temporada", "season"), ("Edad", "age"), ("NAC", "nat"),
            ("POS", "pos"), ("Grupo", "pos_group"), ("Club", "club"), ("Liga", "league"),
            ("PJ Liga", "LIGA_PJ"), ("Min Liga", "LIGA_Min"), ("Min Total", "TOT_Min"),
            ("Score en su liga", "score_league"), ("Score global posición", "score_global"),
            ("Cobertura datos", "pos_coverage"), ("Pool liga pequeño", "league_pool_small"),
            ("Calidad datos liga", "competition_data_quality"),
            ("COMPETITION_STRENGTH", "competition_strength"), ("OPPONENT_STRENGTH", "opponent_strength"),
            ("CONTEXT_SCORE", "context_score"), ("TEAM_STRENGTH", "team_strength"),
            ("REL_PERF", "REL_PERF"), ("Tendencia CA", "trend_ca"),
            ("ELO", "ELO"), ("FORM", "FORM"), ("CONSISTENCY", "CONSISTENCY"),
            ("Rol principal", "best_role"), ("Nota rol", "best_role_score"), ("Aviso", "league_label_warning")]
    formula_cols = ["CONF_MUESTRA", "BASE_LIGA", "CA_CONTEXT", "CA_RAW", "CA_FINAL", "CA_CONFIDENCE",
                    "CRECIMIENTO", "PA_ESTIMATE", "PA_RANGE_LOW", "PA_RANGE_HIGH", "PA_CONFIDENCE", "AGE_SCORE", "SCOUT_SCORE", "model_version"]
    df = pd.DataFrame({h: b[c] if c in b else np.nan for h, c in cols})
    for fc in formula_cols:
        df[fc] = None
    write_df(ws, df, freeze="D2")
    C = {h: get_column_letter(j) for j, h in enumerate(df.columns, 1)}
    n = len(df)
    for i in range(2, n + 2):
        g = lambda h: f"{C[h]}{i}"
        mins, sl, sg, ctx = g("Min Liga"), g("Score en su liga"), g("Score global posición"), g("CONTEXT_SCORE")
        f = {}
        f["CONF_MUESTRA"] = f'=IF(AND(ISNUMBER({mins}),{mins}>={P["min_minutos_valorar"]},ISNUMBER({ctx}),ISNUMBER({sl})),1-EXP(-{mins}/{P["k_min_confianza"]}),"")'
        cm = g("CONF_MUESTRA")
        f["BASE_LIGA"] = f'=IF(ISNUMBER({cm}),{P["factor_base_liga"]}*{ctx},"")'
        bl = g("BASE_LIGA")
        f["CA_CONTEXT"] = f'=IF(ISNUMBER({cm}),MAX(0,MIN(100,{bl}+{P["pendiente_contexto"]}*({sl}-50))),"")'
        f["CA_RAW"] = f'=IF(ISNUMBER({cm}),MAX(0,MIN(100,{P["centro_raw"]}+{P["pendiente_raw"]}*({sg}-50))),"")'
        cc, cr = g("CA_CONTEXT"), g("CA_RAW")
        f["CA_FINAL"] = f'=IF(ISNUMBER({cm}),MAX(0,MIN(100,{cm}*({P["peso_contexto"]}*{cc}+{P["peso_raw"]}*{cr})+(1-{cm})*{bl})),"")'
        ca = g("CA_FINAL")
        f["CA_CONFIDENCE"] = (f'=IF(ISNUMBER({cm}),100*({P["conf_peso_muestra"]}*{cm}+{P["conf_peso_cobertura"]}*N({g("Cobertura datos")})'
                              f'+{P["conf_peso_calidad_liga"]}*N({g("Calidad datos liga")})*IF({g("Pool liga pequeño")}=TRUE,0.5,1)),"")')
        age, rp, tr = g("Edad"), g("REL_PERF"), g("Tendencia CA")
        tab = P["tabla_edad"]
        f["CRECIMIENTO"] = (f'=IF(AND(ISNUMBER({ca}),ISNUMBER({age}),ISNUMBER({rp})),MAX(0,VLOOKUP(MAX(15,MIN(45,{age})),{tab},2,TRUE)'
                            f'*({P["pa_mult_min"]}+({P["pa_mult_max"]}-{P["pa_mult_min"]})*{rp}/100)*MAX(0,MIN(1,(100-{ca})/{P["pa_margen_ref"]}))'
                            f'+{P["pa_peso_tendencia"]}*N({tr})),"")')
        gr = g("CRECIMIENTO")
        f["PA_ESTIMATE"] = f'=IF(ISNUMBER({gr}),MIN(99,{ca}+{gr}),"")'
        pa, cf = g("PA_ESTIMATE"), g("CA_CONFIDENCE")
        width = f'VLOOKUP(MAX(15,MIN(45,{age})),{tab},3,TRUE)*(1.5-{cf}/100)'
        f["PA_RANGE_LOW"] = f'=IF(ISNUMBER({pa}),MAX({ca},{pa}-{width}),"")'
        f["PA_RANGE_HIGH"] = f'=IF(ISNUMBER({pa}),MIN(100,{pa}+{width}),"")'
        f["PA_CONFIDENCE"] = f'=IF(ISNUMBER({pa}),{cf}*VLOOKUP(MAX(15,MIN(45,{age})),{tab},4,TRUE),"")'
        f["AGE_SCORE"] = f'=IF(ISNUMBER({age}),VLOOKUP(MAX(15,MIN(45,{age})),{tab},5,TRUE),"")'
        ags, fo = g("AGE_SCORE"), g("FORM")
        w = {k: P[f"scout_{k}"] for k in M["scout_score"]["weights"]}
        comps = [("PA_ESTIMATE", pa), ("CA_FINAL", ca), ("REL_PERF", rp), ("AGE_SCORE", ags), ("FORM", fo), ("CONFIDENCE", cf)]
        numr = "+".join(f"IF(ISNUMBER({ref}),{w[k]}*{ref},0)" for k, ref in comps)
        denr = "+".join(f"IF(ISNUMBER({ref}),{w[k]},0)" for k, ref in comps)
        f["SCOUT_SCORE"] = f'=IF(ISNUMBER({ca}),({numr})/({denr}),"")'
        f["model_version"] = f'={P["model_version"]}'
        for h, fx in f.items():
            c = ws[g(h)]
            c.value, c.font = fx, F_BASE
    for h in formula_cols:
        ws[f"{C[h]}1"].fill = PatternFill("solid", fgColor="375623")
    fmt = {"Score en su liga": "0.0", "Score global posición": "0.0", "Cobertura datos": "0%", "Calidad datos liga": "0.00",
           "COMPETITION_STRENGTH": "0", "OPPONENT_STRENGTH": "0", "CONTEXT_SCORE": "0.0", "TEAM_STRENGTH": "0.0", "REL_PERF": "0.0",
           "Tendencia CA": "+0.0;-0.0;0.0", "ELO": "0.0", "FORM": "0.0", "CONSISTENCY": "0.0", "Nota rol": "0.0",
           "CONF_MUESTRA": "0%", "BASE_LIGA": "0.0", "CA_CONTEXT": "0.0", "CA_RAW": "0.0", "CA_FINAL": "0.0", "CA_CONFIDENCE": "0",
           "CRECIMIENTO": "0.0", "PA_ESTIMATE": "0.0", "PA_RANGE_LOW": "0.0", "PA_RANGE_HIGH": "0.0", "PA_CONFIDENCE": "0", "AGE_SCORE": "0", "SCOUT_SCORE": "0.0"}
    for h, nf in fmt.items():
        for i in range(2, n + 2):
            ws[f"{C[h]}{i}"].number_format = nf
    ws.column_dimensions[C["KEY"]].hidden = True
    ws.column_dimensions[C["Jugador"]].width = 24
    ws.column_dimensions[C["Club"]].width = 20
    ws.column_dimensions[C["Liga"]].width = 16
    for h in ["CA_FINAL", "PA_ESTIMATE", "SCOUT_SCORE"]:
        scale(ws, f"{C[h]}2:{C[h]}{n + 1}")
    for h in ["OPPONENT_STRENGTH", "ELO", "FORM", "CONSISTENCY", "REL_PERF", "Tendencia CA", "Rol principal", "Nota rol", "Edad"]:
        grey_blanks(ws, f"{C[h]}2:{C[h]}{n + 1}")
    val_cols = C

    # ---------------- ATRIBUTOS (misma fila que VALORACION)
    wa = wb.create_sheet("ATRIBUTOS")
    ad = pd.DataFrame({"KEY": b.KEY, "Jugador": b.name, "Grupo": b.pos_group, "Club": b.club, "Liga": b.league})
    for a, d in attrs.items():
        ad[d["label"]] = b[f"attrL_{a}"].round(1)
    for a, d in attrs.items():
        ad[f"cob. {d['label']}"] = b[f"cov_{a}"].round(2)
    write_df(wa, ad, freeze="C2")
    na = len(attrs)
    rng = f"F2:{get_column_letter(5 + na)}{len(ad) + 1}"
    scale(wa, rng)
    grey_blanks(wa, rng)
    wa.column_dimensions["A"].hidden = True
    wa.cell(row=1, column=5 + na + 1).comment = Comment("Cobertura = parte del peso del atributo que tiene dato. < 50% ⇒ atributo UNKNOWN (gris).", "modelo")

    # ---------------- ROLES
    wr = wb.create_sheet("ROLES")
    rd = pd.DataFrame({"KEY": b.KEY, "Jugador": b.name, "Grupo": b.pos_group, "Club": b.club, "Liga": b.league})
    for r, d in roles.items():
        rd[f"{d['label']} [{'/'.join(d['groups'])}]"] = b[f"role_{r}"].round(1)
    for r, d in roles.items():
        rd[f"cob. {d['label']}"] = b[f"rolecov_{r}"].round(2)
    write_df(wr, rd, freeze="C2")
    rng = f"F2:{get_column_letter(5 + len(roles))}{len(rd) + 1}"
    scale(wr, rng)
    grey_blanks(wr, rng)
    wr.column_dimensions["A"].hidden = True

    # ---------------- PERCENTILES (métricas con algún dato)
    wp = wb.create_sheet("PERCENTILES")
    pm = [m for m in metrics if f"pctL_{m}" in b and b[f"pctL_{m}"].notna().any()]
    pdf = pd.DataFrame({"KEY": b.KEY, "Jugador": b.name, "Grupo": b.pos_group, "Club": b.club, "Liga": b.league, "Min Liga": b.LIGA_Min})
    for m in pm:
        lab = metrics[m]["label"]
        pdf[f"{lab} · valor"] = b[m].round(3)
        if metrics[m].get("shrink"):
            pdf[f"{lab} · ajust. muestra"] = b[f"adj_{m}"].round(3)
        pdf[f"{lab} · pct liga"] = b[f"pctL_{m}"].round(1)
        pdf[f"{lab} · pct global"] = b[f"pctG_{m}"].round(1)
    write_df(wp, pdf, freeze="C2")
    for j, col in enumerate(pdf.columns, 1):
        if "pct" in col:
            scale(wp, f"{get_column_letter(j)}2:{get_column_letter(j)}{len(pdf) + 1}")
    grey_blanks(wp, f"G2:{get_column_letter(len(pdf.columns))}{len(pdf) + 1}")
    wp.column_dimensions["A"].hidden = True

    # ---------------- DESGLOSE_CA (de la nota a la estadística)
    wd = wb.create_sheet("DESGLOSE_CA")
    groups = cfg["positions"]["groups"]
    rated = b[b.CA_FINAL.notna()]
    chunks = []
    for g, gd in groups.items():
        sub = rated[rated.pos_group == g]
        if not len(sub):
            continue
        gw = gd["ca_weights"]
        tot = sum(w * sub[f"attrL_{a}"].notna() for a, w in gw.items())
        for a, w in gw.items():
            av = sub[f"attrL_{a}"]
            ok = av.notna()
            for mname, mw in attrs[a]["components"].items():
                pl = sub[f"pctL_{mname}"] if f"pctL_{mname}" in sub else pd.Series(np.nan, index=sub.index)
                m2 = ok & pl.notna()
                if not m2.any():
                    continue
                s2 = sub[m2]
                chunks.append(pd.DataFrame({
                    "KEY": s2.KEY, "Jugador": s2["name"], "Temporada": s2.season, "Club": s2.club, "Grupo": g,
                    "CA_FINAL": s2.CA_FINAL.round(1), "Score en su liga": s2.score_league.round(1),
                    "Atributo": attrs[a]["label"], "Peso atributo (config)": w,
                    "Peso efectivo": (w / tot[m2]).round(3), "Valor atributo (liga)": av[m2].round(1),
                    "Aporte al score": (av[m2] * w / tot[m2]).round(2),
                    "Estadística": metrics[mname]["label"], "Peso en atributo": mw,
                    "Valor bruto (per90/ratio)": s2[mname].astype(float).round(3),
                    "Ajustado por muestra": s2[f"adj_{mname}"].astype(float).round(3),
                    "Pct liga": pl[m2].round(1), "Pct global": s2[f"pctG_{mname}"].round(1)}))
    dd = pd.concat(chunks, ignore_index=True).sort_values(["Temporada", "Jugador", "Aporte al score"], ascending=[True, True, False]) if chunks else pd.DataFrame()
    write_df(wd, dd, freeze="C2")
    wd.column_dimensions["A"].hidden = True
    wd.cell(row=1, column=len(dd.columns) + 2, value="Solo filas con dato. Los atributos UNKNOWN no aportan y reducen la 'Cobertura datos' (ver ATRIBUTOS y CONFIG_PESOS).").font = F_GREY

    # ---------------- CONTEXTO: competiciones y equipos
    wc = wb.create_sheet("COMPETICIONES")
    comps = cfg.get("competitions", {}).get("leagues", {})
    obs = b.groupby("league").agg(jugadores=("player_id", "count"), con_CA=("CA_FINAL", "count"),
                                  clubes=("club", "nunique")).reset_index()
    cd = pd.DataFrame([{"Liga (tal cual en tu Excel)": k, "País": v.get("country"), "Confederación": v.get("confederation"),
                        "COMPETITION_STRENGTH (propio)": v.get("strength"), "Calidad datos (0-1)": v.get("data_quality"),
                        "Nota": v.get("note", "")} for k, v in comps.items()])
    if len(cd):
        cd = cd.merge(obs, left_on="Liga (tal cual en tu Excel)", right_on="league", how="left").drop(columns=["league"])
    write_df(wc, cd, freeze="B2")
    wc["A" + str(len(cd) + 3)] = "competition_strength es una valoración PROPIA del modelo (no ranking oficial FIFA/UEFA). Se edita en config/competitions.json y puede cambiar por temporada."
    wc["A" + str(len(cd) + 3)].font = F_GREY
    bl = cfg["model"]["rating_blocks"]
    r0 = len(cd) + 5
    hdr(wc, r0, ["Bloque", "Categoría", "counts_for_statistics", "counts_for_rating", "counts_for_elo", "counts_for_form"])
    for i, (k, v) in enumerate(bl.items(), r0 + 1):
        for j, x in enumerate([k, v["category"], v["counts_for_statistics"], v["counts_for_rating"], v["counts_for_elo"], v["counts_for_form"]], 1):
            wc.cell(row=i, column=j, value=x).font = F_BASE

    we = wb.create_sheet("EQUIPOS")
    tcols = ["season", "league", "club", "team_matches_proxy", "team_gf_p90_proxy", "team_gc_p90", "team_gd_p90_proxy", "team_strength_in_league", "team_strength"]
    t = base.drop_duplicates(["season", "league", "club"])[tcols].copy()
    t["team_gc_p90"] = base[base.pos_group != "GK"].drop_duplicates(["season", "league", "club"]).set_index(["season", "league", "club"]).team_gc_p90.reindex(pd.MultiIndex.from_frame(t[["season", "league", "club"]])).values
    t.columns = ["Temporada", "Liga", "Club", "PJ equipo (proxy)", "GF/partido (proxy)", "GC/90 (porteros)", "Dif./partido (proxy)", "Fuerza en su liga (pct)", "TEAM_STRENGTH"]
    write_df(we, t.sort_values(["Liga", "Fuerza en su liga (pct)"], ascending=[True, False]).round(2), freeze="D2")

    # ---------------- ELO_PARTIDOS
    wm = wb.create_sheet("ELO_PARTIDOS")
    mcols = ["name", "season", "date", "block", "club", "opponent", "home", "gf", "ga", "Min", "G", "A", "GC", "sofa_rating", "match_rating",
             "opp_elo", "elo_before", "expected", "score", "elo_delta"]
    if matches is not None and len(matches):
        mx = matches.copy()
        if "name" not in mx:
            mx["name"] = mx.player_id.map(b.drop_duplicates("player_id").set_index("player_id").name)
        mx["date"] = pd.to_datetime(mx.date).dt.strftime("%Y-%m-%d")
        for c in ("match_rating", "opp_elo", "elo_before", "expected", "score", "elo_delta"):
            if c in mx:
                mx[c] = pd.to_numeric(mx[c], errors="coerce").round(3 if c in ("expected", "score") else 1)
        mx = mx.sort_values(["name", "date"])
        write_df(wm, mx[[c for c in mcols if c in mx]], freeze="C2")
    else:
        hdr(wm, 1, mcols)
        wm["A3"] = "SIN DATOS POR PARTIDO EN ESTA PRUEBA → ELO, FORM y CONSISTENCY = UNKNOWN (no se inventan ni se derivan del CA)."
        wm["A3"].font = F_BOLD
        wm["A4"] = "El motor ya está programado (engine.match_layer): Elo de equipos a partir de resultados → fuerza del rival → ELO del jugador por partido (K × minutos × importancia) → FORM_5/10/20 con decaimiento → CONSISTENCY por desviación ponderada."
        wm["A5"] = "Para activarlo necesitamos una fila por jugador y partido (fecha, rival, local/visitante, resultado, minutos, G, A). Transfermarkt 'performance-game' la da."
        for c in ("A4", "A5"):
            wm[c].font = F_GREY

    # ---------------- SIMILARES
    wsim = wb.create_sheet("SIMILARES")
    if sims is not None and len(sims):
        s = sims.merge(b[["player_id", "season", "KEY"]], on=["player_id", "season"], how="left")
        s = s[["KEY", "name", "pos_group", "rank", "similar_name", "similar_club", "similar_league", "similarity", "metrics_used"]]
        s.columns = ["KEY", "Jugador", "Grupo", "Nº", "Similar", "Club similar", "Liga similar", "Similitud (0-100)", "Métricas usadas"]
        write_df(wsim, s, freeze="C2")
        wsim.column_dimensions["A"].hidden = True
    # ---------------- BUSCADOR_TALENTO (listas prefiltradas)
    wt = wb.create_sheet("BUSCADOR_TALENTO")
    wt["A1"] = "BUSCADOR DE TALENTO — señales para investigar, NO conclusiones"
    wt["A1"].font = F_TITLE
    filt = [
        ("Jóvenes ≤21 con REL_PERF ≥ 85 y confianza ≥ 45", (b.age <= 21) & (b.REL_PERF >= 85) & (b.CA_CONFIDENCE >= 45)),
        ("Jóvenes ≤21 destacando en ligas de COMPETITION_STRENGTH < 80", (b.age <= 21) & (b.REL_PERF >= 80) & (b.competition_strength < 80) & (b.CA_CONFIDENCE >= 40)),
        ("≤23 con PA estimado ≥ 85 (confianza PA ≥ 30)", (b.age <= 23) & (b.PA_ESTIMATE >= 85) & (b.PA_CONFIDENCE >= 30)),
    ]
    tc = ["name", "age", "pos", "club", "league", "LIGA_Min", "CA_FINAL", "CA_CONFIDENCE", "PA_ESTIMATE", "PA_RANGE_LOW", "PA_RANGE_HIGH", "PA_CONFIDENCE", "REL_PERF", "SCOUT_SCORE", "competition_strength"]
    th = ["Jugador", "Edad", "POS", "Club", "Liga", "Min Liga", "CA", "CA conf.", "PA", "PA bajo", "PA alto", "PA conf.", "REL_PERF", "SCOUT", "Fuerza liga"]
    r = 3
    for title, mask in filt:
        wt.cell(row=r, column=1, value=title).font = F_BOLD
        sub = b[mask].sort_values("SCOUT_SCORE", ascending=False).head(40)[tc].round(1)
        sub.columns = th
        hdr(wt, r + 1, th)
        for i, row in enumerate(sub.astype(object).where(sub.notna(), None).values.tolist(), r + 2):
            for j, v in enumerate(row, 1):
                wt.cell(row=i, column=j, value=v).font = F_BASE
        r += len(sub) + 4
    wt.cell(row=r, column=1, value="Valores calculados por el script (versión del modelo en PARAMETROS). Para otros cortes usa el autofiltro de VALORACION.").font = F_GREY
    for j, w in enumerate([24, 6, 6, 20, 16, 9, 7, 8, 7, 8, 8, 8, 9, 8, 10], 1):
        wt.column_dimensions[get_column_letter(j)].width = w

    # ---------------- HISTORIAL
    if history is not None and len(history):
        wh = wb.create_sheet("HISTORIAL")
        write_df(wh, history, freeze="C2")

    # ---------------- CORRELACIONES
    if corr is not None and len(corr):
        wco = wb.create_sheet("CORRELACIONES")
        write_df(wco, corr, freeze="A2")

    # ---------------- CONFIG (espejo legible de los JSON)
    wcfg = wb.create_sheet("CONFIG_PESOS")
    r = 1
    wcfg.cell(row=r, column=1, value="ATRIBUTOS: estadísticas y pesos (config/attributes.json)").font = F_TITLE
    r += 1
    hdr(wcfg, r, ["Atributo", "Familia", "Estadística", "Peso", "Disponibilidad", "Grupo de correlación", "Dirección"])
    for a, d in attrs.items():
        for m, w in d["components"].items():
            r += 1
            for j, v in enumerate([d["label"], d["family"], metrics[m]["label"], w, metrics[m]["availability"], metrics[m].get("corr_group"), metrics[m]["direction"]], 1):
                wcfg.cell(row=r, column=j, value=v).font = F_BASE
    r += 3
    wcfg.cell(row=r, column=1, value="PESO DE CADA ATRIBUTO EN EL CA POR POSICIÓN (config/positions.json)").font = F_TITLE
    r += 1
    gnames = list(groups)
    hdr(wcfg, r, ["Atributo"] + [f"{g} · {groups[g]['label']}" for g in gnames])
    for a, d in attrs.items():
        r += 1
        wcfg.cell(row=r, column=1, value=d["label"]).font = F_BASE
        for j, g in enumerate(gnames, 2):
            v = groups[g]["ca_weights"].get(a)
            if v:
                wcfg.cell(row=r, column=j, value=v).font = F_BASE
    r += 3
    wcfg.cell(row=r, column=1, value="ROLES (config/roles.json)").font = F_TITLE
    r += 1
    hdr(wcfg, r, ["Rol", "Grupos", "Atributo", "Peso"])
    for k, d in roles.items():
        for a, w in d["weights"].items():
            r += 1
            for j, v in enumerate([d["label"], "/".join(d["groups"]), attrs[a]["label"], w], 1):
                wcfg.cell(row=r, column=j, value=v).font = F_BASE
    for j, w in enumerate([34, 18, 40, 10, 14, 18, 18], 1):
        wcfg.column_dimensions[get_column_letter(j)].width = w

    # ---------------- DATOS_ORIGEN
    if raw is not None:
        wo = wb.create_sheet("DATOS_ORIGEN")
        write_df(wo, raw, freeze="C2")

    # ---------------- FICHA (fórmulas)
    build_ficha(ficha, b, val_cols, attrs, roles, pm, metrics, n)
    build_comparador(comp, b, val_cols, attrs, n)
    build_leeme(leeme, cfg, b, source_note)
    wb.save(out_path)
    return b


def _lk(sheet, col_letter, rowref):
    return f'IFERROR(INDEX({sheet}!${col_letter}:${col_letter},{rowref}),"")'


def build_ficha(ws, b, C, attrs, roles, pm, metrics, n):
    ws["A1"] = "FICHA DE JUGADOR"
    ws["A1"].font = F_TITLE
    ws["A3"] = "Jugador (elige de la lista o escribe 'Nombre | Club | Temporada'):"
    ws["A3"].font = F_BOLD
    top = b.sort_values("CA_FINAL", ascending=False).KEY.iloc[0]
    ws["B4"] = top
    ws["B4"].font, ws["B4"].fill = Font(name=FONT, size=12, bold=True, color="0000FF"), FILL_INPUT
    ws.merge_cells("B4:E4")
    dv = DataValidation(type="list", formula1=f"=VALORACION!$A$2:$A${n + 1}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    ws["A5"] = "Fila:"
    ws["B5"] = '=IFERROR(MATCH(B4,VALORACION!$A:$A,0),"no encontrado")'
    ws["A5"].font = F_GREY
    R = "$B$5"
    def put(r, c, label, formula, fmt=None):
        ws.cell(row=r, column=c, value=label).font = F_BOLD
        cell = ws.cell(row=r, column=c + 1, value=formula)
        cell.font = F_BASE
        if fmt:
            cell.number_format = fmt
    info = [("Jugador", "Jugador"), ("Edad", "Edad"), ("POS", "POS"), ("Club", "Club"), ("Liga", "Liga"), ("Temporada", "Temporada"),
            ("Min Liga", "Min Liga"), ("PJ Liga", "PJ Liga"), ("Min Total", "Min Total")]
    r = 7
    ws.cell(row=r, column=1, value="JUGADOR").font = F_HDR
    ws.cell(row=r, column=1).fill = FILL_HDR
    for i, (lab, h) in enumerate(info, r + 1):
        put(i, 1, lab, "=" + _lk("VALORACION", C[h], R))
    ratings = [("CURRENT ABILITY", "CA_FINAL", "0.0"), ("CA confidence", "CA_CONFIDENCE", "0\"%\""), ("  CA_CONTEXT", "CA_CONTEXT", "0.0"),
               ("  CA_RAW", "CA_RAW", "0.0"), ("POTENTIAL (estim.)", "PA_ESTIMATE", "0.0"), ("  PA rango bajo", "PA_RANGE_LOW", "0.0"),
               ("  PA rango alto", "PA_RANGE_HIGH", "0.0"), ("  PA confidence", "PA_CONFIDENCE", "0\"%\""), ("ELO", "ELO", "0.0"),
               ("FORM", "FORM", "0.0"), ("CONSISTENCY", "CONSISTENCY", "0.0"), ("SCOUT SCORE", "SCOUT_SCORE", "0.0"),
               ("REL_PERF", "REL_PERF", "0.0"), ("Rol principal", "Rol principal", None), ("Nota rol", "Nota rol", "0.0")]
    ws.cell(row=r, column=4, value="VALORACIONES").font = F_HDR
    ws.cell(row=r, column=4).fill = FILL_HDR
    for i, (lab, h, fm) in enumerate(ratings, r + 1):
        put(i, 4, lab, f'=IF({_lk("VALORACION", C[h], R)}="","UNKNOWN",{_lk("VALORACION", C[h], R)})', fm)
    ctx = [("COMPETITION_STRENGTH", "COMPETITION_STRENGTH"), ("OPPONENT_STRENGTH", "OPPONENT_STRENGTH"), ("CONTEXT_SCORE", "CONTEXT_SCORE"),
           ("TEAM_STRENGTH", "TEAM_STRENGTH"), ("Confianza muestra", "CONF_MUESTRA"), ("Cobertura datos", "Cobertura datos"), ("Versión modelo", "model_version"), ("Aviso", "Aviso")]
    ws.cell(row=r, column=7, value="CONTEXTO").font = F_HDR
    ws.cell(row=r, column=7).fill = FILL_HDR
    for i, (lab, h) in enumerate(ctx, r + 1):
        put(i, 7, lab, f'=IF({_lk("VALORACION", C[h], R)}="","{"—" if h == "Aviso" else "UNKNOWN"}",{_lk("VALORACION", C[h], R)})',
            "0%" if h in ("CONF_MUESTRA", "Cobertura datos") else "0.0")
    # atributos
    r = 25
    ws.cell(row=r, column=1, value="ATRIBUTOS (0-100, percentil ponderado en su posición y liga)").font = F_HDR
    ws.cell(row=r, column=1).fill = FILL_HDR
    ws.cell(row=r, column=4, value="ROLES").font = F_HDR
    ws.cell(row=r, column=4).fill = FILL_HDR
    for i, (a, d) in enumerate(attrs.items(), r + 1):
        col = get_column_letter(6 + i - r - 1)
        put(i, 1, d["label"], f'=IF({_lk("ATRIBUTOS", col, R)}="","UNKNOWN",{_lk("ATRIBUTOS", col, R)})', "0.0")
    for i, (k, d) in enumerate(roles.items(), r + 1):
        col = get_column_letter(6 + i - r - 1)
        put(i, 4, f"{d['label']} [{'/'.join(d['groups'])}]", f'=IF({_lk("ROLES", col, R)}="","—",{_lk("ROLES", col, R)})', "0.0")
    # percentiles (primeras métricas)
    ws.cell(row=r, column=7, value="ESTADÍSTICAS (valor · pct liga · pct global)").font = F_HDR
    ws.cell(row=r, column=7).fill = FILL_HDR
    # localizar columnas en PERCENTILES
    cols = ["KEY", "Jugador", "Grupo", "Club", "Liga", "Min Liga"]
    for m in pm:
        lab = metrics[m]["label"]
        cols += [f"{lab} · valor"] + ([f"{lab} · ajust. muestra"] if metrics[m].get("shrink") else []) + [f"{lab} · pct liga", f"{lab} · pct global"]
    pos = {c: get_column_letter(j) for j, c in enumerate(cols, 1)}
    for i, m in enumerate(pm, r + 1):
        lab = metrics[m]["label"]
        ws.cell(row=i, column=7, value=lab).font = F_BOLD
        for j, suf in enumerate(["valor", "pct liga", "pct global"], 8):
            c = ws.cell(row=i, column=j, value=f'=IF({_lk("PERCENTILES", pos[f"{lab} · {suf}"], R)}="","UNKNOWN",{_lk("PERCENTILES", pos[f"{lab} · {suf}"], R)})')
            c.font = F_BASE
            c.number_format = "0.00" if suf == "valor" else "0"
    # similares
    rs = r + len(pm) + 3
    ws.cell(row=rs, column=7, value="JUGADORES SIMILARES (perfil básico de percentiles: orientativo)").font = F_HDR
    ws.cell(row=rs, column=7).fill = FILL_HDR
    for k in range(5):
        rr = f"MATCH($B$4,SIMILARES!$A:$A,0)+{k}"
        c = ws.cell(row=rs + 1 + k, column=7, value=f'=IFERROR(IF(INDEX(SIMILARES!$A:$A,{rr})=$B$4,INDEX(SIMILARES!$E:$E,{rr})&" ("&INDEX(SIMILARES!$F:$F,{rr})&")",""),"")')
        c.font = F_BASE
        c = ws.cell(row=rs + 1 + k, column=8, value=f'=IFERROR(IF(INDEX(SIMILARES!$A:$A,{rr})=$B$4,INDEX(SIMILARES!$H:$H,{rr}),""),"")')
        c.font = F_BASE
    for rng_ in ["E8:E30", "B26:B60", "E26:E60", "I26:J40"]:
        ws.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=['"UNKNOWN"'], fill=FILL_UNK, font=F_GREY))
    ws.conditional_formatting.add("B26:B60", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))
    for col, w in zip("ABCDEFGHIJ", [30, 16, 3, 34, 12, 3, 34, 11, 11, 11]):
        ws.column_dimensions[col].width = w
    ws["A62"] = "Gris 'UNKNOWN' = no hay dato para calcularlo (nunca se rellena con 0). Para ver de qué estadísticas sale cada número: hoja DESGLOSE_CA, filtra por el jugador."
    ws["A62"].font = F_GREY


def build_comparador(ws, b, C, attrs, n):
    ws["A1"] = "COMPARADOR 1 vs 1"
    ws["A1"].font = F_TITLE
    keys = b.sort_values("CA_FINAL", ascending=False).KEY.tolist()
    ws["A3"], ws["A4"] = "Jugador A", "Jugador B"
    for c in ("A3", "A4"):
        ws[c].font = F_BOLD
    ws["B3"], ws["B4"] = keys[0], keys[1] if len(keys) > 1 else keys[0]
    dv = DataValidation(type="list", formula1=f"=VALORACION!$A$2:$A${n + 1}")
    ws.add_data_validation(dv)
    for c in ("B3", "B4"):
        ws[c].font, ws[c].fill = Font(name=FONT, size=11, bold=True, color="0000FF"), FILL_INPUT
        dv.add(c)
    ws["F3"] = '=IFERROR(MATCH(B3,VALORACION!$A:$A,0),"")'
    ws["F4"] = '=IFERROR(MATCH(B4,VALORACION!$A:$A,0),"")'
    hdr(ws, 6, ["Métrica", "A", "B", "Diferencia A−B", "Nota"])
    items = [(h, "VALORACION", C[h]) for h in ["Edad", "Club", "Liga", "Min Liga", "CA_FINAL", "CA_CONFIDENCE", "CA_CONTEXT", "CA_RAW",
                                                "PA_ESTIMATE", "PA_CONFIDENCE", "ELO", "FORM", "CONSISTENCY", "REL_PERF", "SCOUT_SCORE",
                                                "COMPETITION_STRENGTH", "TEAM_STRENGTH", "Rol principal"]]
    items += [(d["label"], "ATRIBUTOS", get_column_letter(6 + i)) for i, (a, d) in enumerate(attrs.items())]
    for i, (lab, sh, col) in enumerate(items, 7):
        ws.cell(row=i, column=1, value=lab).font = F_BOLD
        for j, rr in [(2, "$F$3"), (3, "$F$4")]:
            ws.cell(row=i, column=j, value=f'=IF({_lk(sh, col, rr)}="","UNKNOWN",{_lk(sh, col, rr)})').font = F_BASE
            ws.cell(row=i, column=j).number_format = "0.0"
        c = ws.cell(row=i, column=4, value=f'=IF(AND(ISNUMBER(B{i}),ISNUMBER(C{i})),B{i}-C{i},"")')
        c.font, c.number_format = F_BASE, "+0.0;-0.0;0.0"
        if lab not in ("Edad", "Min Liga", "Club", "Liga", "Rol principal"):
            ws.cell(row=i, column=5, value=f'=IF(D{i}="","",IF(ABS(D{i})<3,"similar",IF(D{i}>0,"más alto A","más alto B")))').font = F_GREY
    last = 6 + len(items)
    ws.conditional_formatting.add(f"B7:C{last}", CellIsRule(operator="equal", formula=['"UNKNOWN"'], fill=FILL_UNK, font=F_GREY))
    ws.conditional_formatting.add(f"D7:D{last}", ColorScaleRule(start_type="num", start_value=-20, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=20, end_color="63BE7B"))
    for col, w in zip("ABCDEF", [30, 26, 26, 16, 14, 6]):
        ws.column_dimensions[col].width = w
    ws["A" + str(last + 2)] = "Diferencias < 3 puntos = 'similar': dentro del margen de ruido del modelo. Mira siempre la confianza de ambos."
    ws["A" + str(last + 2)].font = F_GREY


def build_leeme(ws, cfg, b, source_note):
    M = cfg["model"]
    rated = b.CA_FINAL.notna().sum()
    lines = [
        ("SISTEMA DE VALORACIÓN FC — PRUEBA " + M["model_version"], F_TITLE),
        (f"Generado {M['model_date']}. {source_note}", F_GREY),
        ("", None),
        ("QUÉ ES", F_BOLD),
        ("Un sistema propio y explicable. NO copia el ELO de BeSoccer ni el CA/PA de Football Manager. Cada número se puede seguir hasta la estadística original.", F_BASE),
        ("", None),
        ("FLUJO", F_BOLD),
        ("DATOS_ORIGEN → per90 → ajuste por muestra → PERCENTILES (posición; liga y global) → ATRIBUTOS → ROLES → CA → REL_PERF → PA → SCOUT.  ELO / FORM / CONSISTENCY: capa por partido.", F_BASE),
        ("", None),
        ("MÉTRICAS (separadas, no son la misma cosa)", F_BOLD),
        ("CA (Current Ability): nivel actual estimado. CA_RAW = sin contexto; CA_CONTEXT = cómo domina en su liga, situado según el nivel de esa liga; CA_FINAL = mezcla, encogida hacia la media de su liga si hay pocos minutos.", F_BASE),
        ("PA (Potential Ability): ESTIMACIÓN con rango y confianza. Depende de edad, rendimiento relativo, margen hasta el techo y tendencia. No es una certeza.", F_BASE),
        ("ELO: rendimiento competitivo partido a partido frente a la fuerza del rival. Independiente del CA (no se calcula a partir del CA).", F_BASE),
        ("FORM: rendimiento reciente (últimos 5/10/20 partidos, los recientes pesan más). CONSISTENCY: regularidad (desviación de las notas de partido).", F_BASE),
        ("REL_PERF: cuánto destaca frente a jugadores de su posición y franja de edad, en su contexto. Es una señal, NO una conversión de liga.", F_BASE),
        ("SCOUT SCORE: interés para ojear (no 'el mejor'). CONFIDENCE: cuánto fiarse (minutos, cobertura de datos, calidad de datos de la liga).", F_BASE),
        ("", None),
        ("CÓMO USARLO", F_BOLD),
        ("FICHA: elige un jugador en la celda amarilla. COMPARADOR: dos jugadores y sus diferencias. VALORACION: tabla completa con autofiltro (edad, liga, POS, confianza…).", F_BASE),
        ("DESGLOSE_CA: filtra por un jugador y verás atributo → peso → estadística → valor bruto → percentil. BUSCADOR_TALENTO: listas prefiltradas.", F_BASE),
        ("PARAMETROS: celdas amarillas editables. CA/PA/SCOUT de VALORACION son fórmulas y se recalculan solas. Percentiles, atributos y roles son valores: se regeneran con el script (auxiliares/valoracion).", F_BASE),
        ("", None),
        ("LIMITACIONES DE ESTA PRUEBA (importante)", F_BOLD),
    ]
    has_adv = any(c in b and b[c].notna().any() for c in ("npxg_p90", "tackles_won_p90", "pass_cmp_pct"))
    has_elo = "ELO" in b and b.ELO.notna().any()
    if has_adv:
        n_adv = int(b.npxg_p90.notna().sum()) if "npxg_p90" in b else 0
        lines += [("1. Métricas avanzadas (npxG, xA, tiros, pases, regates, entradas, intercepciones, paradas…) de Understat y FotMob (Opta) SOLO para las 5 grandes ligas "
                   f"({n_adv} jugadores-temporada). Resto de ligas: UNKNOWN (gris), no 0.", F_BASE),
                  ("2. Siguen UNKNOWN (las fuentes no las dan): duelos aéreos, centros, pérdidas, errores, pases progresivos, conducciones, SCA y toques en el área. FotMob solo lista a quien supera ~9 % de los minutos de liga (porteros ~50 %).", F_BASE),
                  ("3. Fuera de las 5 grandes el CA se apoya en goles/asistencias por 90, minutos, rendimiento continental y goles encajados: su 'Cobertura datos' es menor y la confianza también.", F_BASE)]
    else:
        lines += [("1. Tus Excel tienen PJ, minutos, G, A (y GC/CS en porteros) por competición. NO tienen xG, pases, entradas, regates, etc. Esos atributos salen UNKNOWN (gris), no 0.", F_BASE),
                  ("2. FBref perdió sus estadísticas avanzadas de Opta en enero de 2026. Fuente preparada: auxiliares/avanzadas/descargar_fotmob_understat.py (FotMob + Understat).", F_BASE),
                  ("3. Por eso el CA se apoya en: producción de goles/asistencias por 90, peso en el equipo (minutos), rendimiento continental y, en defensas/porteros, goles encajados. La 'Cobertura datos' dice qué parte del modelo de cada posición tiene dato.", F_BASE)]
    if has_elo:
        lines += [(f"4. ELO, FORM, CONSISTENCY y OPPONENT_STRENGTH: partido a partido de Transfermarkt SOLO en las 5 grandes ({int(b.ELO.notna().sum())} jugadores-temporada). "
                   "Resto: UNKNOWN. La fuerza del rival sale de un Elo de equipos calculado con los resultados (arranca según la liga y se arrastra de 25-26 a 26-27).", F_BASE)]
    else:
        lines += [("4. Sin datos por partido: ELO, FORM, CONSISTENCY y OPPONENT_STRENGTH = UNKNOWN. Descarga preparada: auxiliares/tmapi/descargar_partidos.py.", F_BASE)]
    lines += [
        ("5. % minutos del equipo y la fuerza del equipo son PROXIES calculados con tus propios datos (máx. PJ de un compañero, goles de la plantilla, GC de sus porteros).", F_BASE),
        ("6. Traspasos: tus Excel fusionan clubes del mismo continente en una fila (regla del 25/09), así que no se puede separar Club A / Club B.", F_BASE),
        ("10. SIMILARES usa solo 4-6 métricas básicas (goles, asistencias, minutos, continental): la similitud es orientativa hasta tener estadísticas avanzadas.", F_BASE),
        ("7. competition_strength es una valoración propia, editable, no oficial.", F_BASE),
        ("8. Columna 'Aviso': clubes con menos de 5 jugadores dentro de una liga grande (Coventry/Ipswich/Hull en 'Premier', clubes austriacos en 'Bundesliga', tunecinos en 'Ligue 1'…). Su columna Liga no parece ser la liga donde jugaron esos minutos, así que no se les da CA hasta verificarlo.", F_BASE),
        ("9. En posiciones con poca cobertura de datos (centrales, medios) la nota se acerca a la media en proporción a lo que falta: jugar todos los minutos no basta para llegar a élite sin más datos.", F_BASE),
        ("", None),
        (f"Jugadores en el archivo: {len(b)} · con CA (≥{M['sample']['min_minutes_rated']}' en liga): {rated}", F_GREY),
        ("Leyenda: gris/UNKNOWN = sin dato · amarillo = editable · cabecera verde oscuro = columna con fórmula · verde→rojo = escala 0-100.", F_GREY),
    ]
    for i, (t, f) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        if f:
            c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 160
